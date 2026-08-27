import asyncio
import base64
import csv
import hashlib
import hmac
import io
import json
import math
import logging
import os
import re
import secrets
import sqlite3
import time
import uuid
from contextlib import asynccontextmanager, suppress
from datetime import datetime, timedelta
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

import httpx
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastembed import TextEmbedding
from langchain_core.output_parsers import JsonOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic import BaseModel, Field
from pymilvus import DataType, MilvusClient
from docx import Document
from pypdf import PdfReader

DATA_DIR = Path(os.getenv("DATA_DIR", "/app/data"))
DB = DATA_DIR / "app.db"
VDB = DATA_DIR / "milvus.db"
MODEL = os.getenv("EMBEDDING_MODEL", "BAAI/bge-small-zh-v1.5")
EMBEDDING_BACKEND = os.getenv("EMBEDDING_BACKEND", "hash")
FASTEMBED_CACHE_PATH = os.getenv("FASTEMBED_CACHE_PATH", "/app/model-cache")
REBUILD_VECTOR_INDEX = os.getenv("REBUILD_VECTOR_INDEX", "0").lower() in {"1", "true", "yes"}
COLLECTION = "interview_qa_" + re.sub(r"[^a-z0-9_]+", "_", f"{EMBEDDING_BACKEND}_{MODEL}_parent_child_user_v3".lower().replace("-", "_")).strip("_")[:96]
DIMENSION = 512
DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY", "")
DEEPSEEK_MODEL = os.getenv("DEEPSEEK_MODEL", "deepseek-v4-pro")
DEEPSEEK_BASE_URL = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com")
SHOWDOC_PUSH_CONFIG = Path(os.getenv("SHOWDOC_PUSH_CONFIG", "/run/secrets/showdoc-push.env"))
ADMIN_USER = os.getenv("ADMIN_USER", "admin")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "admin123456")
APP_VERSION = "4.0.0"
SEARCH_TOP_K = 3
PUSH_TIMEZONE = ZoneInfo("Asia/Shanghai")
logger = logging.getLogger(__name__)


class LoginIn(BaseModel):
    username: str
    password: str


class RegisterIn(BaseModel):
    username: str = Field(min_length=3, max_length=32, pattern=r"^[A-Za-z0-9_.-]+$")
    password: str = Field(min_length=8, max_length=128)
    invite_code: str = Field(min_length=10, max_length=100)


class InviteCreateIn(BaseModel):
    note: str = Field(default="", max_length=100)
    expires_in_days: int = Field(default=7, ge=1, le=365)


class UserStatusIn(BaseModel):
    active: bool


class BankIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    description: str = ""


class QuestionIn(BaseModel):
    question: str = Field(min_length=2, max_length=2000)
    answer: str = Field(min_length=1, max_length=30000)
    category: str = "未分类"
    difficulty: str = "中等"
    position: str = "通用"
    keywords: list[str] = Field(default_factory=list)
    tags: list[str] = Field(default_factory=list)
    source: str = "手动录入"
    bank_id: str | None = None


class QuestionPatchIn(BaseModel):
    question: str | None = None
    answer: str | None = None
    category: str | None = None
    difficulty: str | None = None
    position: str | None = None
    keywords: list[str] | None = None
    tags: list[str] | None = None
    source: str | None = None
    bank_id: str | None = None


class SearchIn(BaseModel):
    query: str = Field(min_length=1, max_length=1000)
    mode: Literal["semantic", "keyword", "hybrid"] = "hybrid"
    bank_id: str | None = None
    category: str | None = None
    difficulty: str | None = None
    position: str | None = None
    limit: int = Field(SEARCH_TOP_K, ge=1, le=20)
    semantic_weight: float = Field(0.75, ge=0, le=1)
    keyword_weight: float = Field(0.25, ge=0, le=1)
    min_score: float = Field(0, ge=0, le=1)


class ReviewIn(BaseModel):
    user_id: str = "default"
    rating: Literal["again", "hard", "good", "easy"]


class InterviewStartIn(BaseModel):
    bank_id: str
    user_id: str = "default"


class InterviewAnswerIn(BaseModel):
    question_id: str
    prompt: str
    answer: str
    depth: int = Field(0, ge=0, le=2)


class AskIn(BaseModel):
    query: str = Field(min_length=1, max_length=2000)
    bank_id: str | None = None


class CustomPushIn(BaseModel):
    title: str = Field(min_length=1, max_length=100)
    content: str = Field(min_length=1, max_length=12000)


class RandomQuestionPushIn(BaseModel):
    bank_id: str | None = None
    include_answer: bool = True


class PushSettingsIn(BaseModel):
    enabled: bool = True
    push_time: str = Field(pattern=r"^(?:[01]\d|2[0-3]):[0-5]\d$")
    bank_id: str | None = None
    include_answer: bool = True
    push_url: str | None = Field(default=None, max_length=500)


class UserAISettingsIn(BaseModel):
    provider: Literal["deepseek", "openai-compatible"] = "deepseek"
    api_base: str = Field(min_length=8, max_length=500)
    model: str = Field(min_length=1, max_length=200)
    api_key: str | None = Field(default=None, max_length=1000)
    enabled: bool = True


class CoachProfileIn(BaseModel):
    target_position: str = Field(min_length=1, max_length=120)
    interview_date: str = Field(default="", max_length=10)
    experience_level: str = Field(default="1-3 年", max_length=40)
    daily_minutes: int = Field(default=30, ge=15, le=180)
    jd_text: str = Field(default="", max_length=30000)
    resume_summary: str = Field(default="", max_length=30000)
    project_summary: str = Field(default="", max_length=30000)
    focus_areas: list[str] = Field(default_factory=list)


class CoachTaskIn(BaseModel):
    completed: bool


def db():
    connection = sqlite3.connect(DB)
    connection.row_factory = sqlite3.Row
    return connection


def add_column(connection, table, column, definition):
    columns = [row["name"] for row in connection.execute(f"PRAGMA table_info({table})").fetchall()]
    if column not in columns:
        connection.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def migrate_owned_content(connection):
    bank_columns = {row["name"] for row in connection.execute("PRAGMA table_info(banks)").fetchall()}
    bank_sql_row = connection.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='banks'").fetchone()
    bank_sql = re.sub(r"\s+", " ", (bank_sql_row["sql"] if bank_sql_row else "").lower())
    if "owner_user_id" not in bank_columns or "name text not null unique" in bank_sql:
        connection.execute("DROP TABLE IF EXISTS banks_owner_migration")
        connection.execute(
            """
            CREATE TABLE banks_owner_migration(
              id TEXT PRIMARY KEY,
              name TEXT NOT NULL,
              description TEXT,
              created_at INTEGER NOT NULL,
              owner_user_id TEXT NOT NULL,
              UNIQUE(owner_user_id, name)
            )
            """
        )
        owner_expression = "COALESCE(owner_user_id, ?)" if "owner_user_id" in bank_columns else "?"
        connection.execute(
            f"""
            INSERT INTO banks_owner_migration(id,name,description,created_at,owner_user_id)
            SELECT id,name,description,created_at,{owner_expression} FROM banks
            """,
            (ADMIN_USER,),
        )
        connection.execute("DROP TABLE banks")
        connection.execute("ALTER TABLE banks_owner_migration RENAME TO banks")

    question_columns = {row["name"] for row in connection.execute("PRAGMA table_info(questions)").fetchall()}
    question_sql_row = connection.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='questions'").fetchone()
    question_sql = re.sub(r"\s+", " ", (question_sql_row["sql"] if question_sql_row else "").lower())
    if "owner_user_id" not in question_columns or "content_hash text unique" in question_sql:
        connection.execute("DROP TABLE IF EXISTS questions_owner_migration")
        connection.execute(
            """
            CREATE TABLE questions_owner_migration(
              id TEXT PRIMARY KEY,
              question TEXT,
              answer TEXT,
              category TEXT,
              difficulty TEXT,
              position TEXT,
              keywords TEXT,
              source TEXT,
              content_hash TEXT,
              created_at INTEGER,
              updated_at INTEGER,
              version INTEGER DEFAULT 1,
              bank_id TEXT,
              tags TEXT DEFAULT '[]',
              chunk_index INTEGER DEFAULT 0,
              parent_id TEXT,
              owner_user_id TEXT NOT NULL,
              UNIQUE(owner_user_id, content_hash)
            )
            """
        )
        owner_expression = "COALESCE(q.owner_user_id, b.owner_user_id, ?)" if "owner_user_id" in question_columns else "COALESCE(b.owner_user_id, ?)"
        connection.execute(
            f"""
            INSERT INTO questions_owner_migration(
              id,question,answer,category,difficulty,position,keywords,source,content_hash,
              created_at,updated_at,version,bank_id,tags,chunk_index,parent_id,owner_user_id
            )
            SELECT q.id,q.question,q.answer,q.category,q.difficulty,q.position,q.keywords,q.source,q.content_hash,
              q.created_at,q.updated_at,q.version,q.bank_id,q.tags,q.chunk_index,q.parent_id,{owner_expression}
            FROM questions q LEFT JOIN banks b ON b.id=q.bank_id
            """,
            (ADMIN_USER,),
        )
        connection.execute("DROP TABLE questions")
        connection.execute("ALTER TABLE questions_owner_migration RENAME TO questions")
    connection.execute("CREATE INDEX IF NOT EXISTS idx_banks_owner ON banks(owner_user_id, created_at)")
    connection.execute("CREATE INDEX IF NOT EXISTS idx_questions_owner_bank ON questions(owner_user_id, bank_id, parent_id)")


def ensure_default_bank(connection, user_id):
    row = connection.execute(
        "SELECT id FROM banks WHERE owner_user_id=? ORDER BY CASE WHEN name='默认题库' THEN 0 ELSE 1 END,created_at LIMIT 1",
        (user_id,),
    ).fetchone()
    if row:
        return row["id"]
    bank_id = str(uuid.uuid4())
    connection.execute(
        "INSERT INTO banks(id,name,description,created_at,owner_user_id) VALUES(?,?,?,?,?)",
        (bank_id, "默认题库", "当前账号的默认题库", int(time.time()), user_id),
    )
    return bank_id


def default_bank_id(user_id):
    with db() as connection:
        return ensure_default_bank(connection, user_id)


def ensure_bank_by_name(name, user_id):
    cleaned = str(name or "").strip()
    if not cleaned:
        return default_bank_id(user_id)
    with db() as connection:
        row = connection.execute("SELECT id FROM banks WHERE owner_user_id=? AND name=?", (user_id, cleaned)).fetchone()
        if row:
            return row["id"]
        bank_id = str(uuid.uuid4())
        connection.execute(
            "INSERT INTO banks(id,name,description,created_at,owner_user_id) VALUES(?,?,?,?,?)",
            (bank_id, cleaned, "Markdown auto-created bank", int(time.time()), user_id),
        )
        return bank_id


def owned_bank(connection, bank_id, user_id):
    if not bank_id:
        return None
    return connection.execute(
        "SELECT * FROM banks WHERE id=? AND owner_user_id=?",
        (bank_id, user_id),
    ).fetchone()


def require_owned_bank(bank_id, user_id):
    with db() as connection:
        bank = owned_bank(connection, bank_id, user_id)
    if not bank:
        raise HTTPException(404, "题库不存在")
    return dict(bank)


def hash_password(password):
    iterations = 310000
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, iterations)
    return "$".join(
        [
            "pbkdf2_sha256",
            str(iterations),
            base64.urlsafe_b64encode(salt).decode("ascii"),
            base64.urlsafe_b64encode(digest).decode("ascii"),
        ]
    )


def verify_password(password, encoded):
    try:
        algorithm, iterations, salt, expected = encoded.split("$", 3)
        if algorithm != "pbkdf2_sha256":
            return False
        digest = hashlib.pbkdf2_hmac(
            "sha256",
            password.encode("utf-8"),
            base64.urlsafe_b64decode(salt.encode("ascii")),
            int(iterations),
        )
        return hmac.compare_digest(base64.urlsafe_b64encode(digest).decode("ascii"), expected)
    except (AttributeError, TypeError, ValueError):
        return False


def normalize_invite_code(code):
    return re.sub(r"\s+", "", str(code or "")).upper()


def invite_code_hash(code):
    return hashlib.sha256(normalize_invite_code(code).encode("utf-8")).hexdigest()


def new_invite_code():
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    groups = ["".join(secrets.choice(alphabet) for _ in range(5)) for _ in range(4)]
    return "JW-" + "-".join(groups)


def init_db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with db() as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS users(
              username TEXT PRIMARY KEY,
              password_hash TEXT NOT NULL,
              role TEXT NOT NULL DEFAULT 'user',
              active INTEGER NOT NULL DEFAULT 1,
              invited_by TEXT,
              created_at INTEGER NOT NULL,
              last_login_at INTEGER
            );
            CREATE TABLE IF NOT EXISTS invites(
              id TEXT PRIMARY KEY,
              code_hash TEXT NOT NULL UNIQUE,
              code_prefix TEXT NOT NULL,
              note TEXT,
              created_by TEXT NOT NULL,
              created_at INTEGER NOT NULL,
              expires_at INTEGER NOT NULL,
              used_by TEXT,
              used_at INTEGER,
              revoked_at INTEGER
            );
            CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY, user_id TEXT, created_at INTEGER, expires_at INTEGER);
            CREATE TABLE IF NOT EXISTS banks(
              id TEXT PRIMARY KEY,
              name TEXT NOT NULL,
              description TEXT,
              created_at INTEGER NOT NULL,
              owner_user_id TEXT NOT NULL,
              UNIQUE(owner_user_id, name)
            );
            CREATE TABLE IF NOT EXISTS questions(
              id TEXT PRIMARY KEY,
              question TEXT,
              answer TEXT,
              category TEXT,
              difficulty TEXT,
              position TEXT,
              keywords TEXT,
              source TEXT,
              content_hash TEXT,
              created_at INTEGER,
              updated_at INTEGER,
              version INTEGER DEFAULT 1,
              bank_id TEXT,
              tags TEXT DEFAULT '[]',
              chunk_index INTEGER DEFAULT 0,
              parent_id TEXT,
              owner_user_id TEXT NOT NULL,
              UNIQUE(owner_user_id, content_hash)
            );
            CREATE TABLE IF NOT EXISTS user_question_states(
              user_id TEXT,
              question_id TEXT,
              mastery_level REAL DEFAULT 0,
              review_count INTEGER DEFAULT 0,
              last_reviewed_at INTEGER,
              next_review_at INTEGER,
              interval_days INTEGER DEFAULT 0,
              ease_factor REAL DEFAULT 2.5,
              stability REAL DEFAULT 1.0,
              difficulty_factor REAL DEFAULT 5.0,
              lapse_count INTEGER DEFAULT 0,
              last_rating TEXT,
              PRIMARY KEY(user_id, question_id)
            );
            CREATE TABLE IF NOT EXISTS search_logs(id INTEGER PRIMARY KEY AUTOINCREMENT, query TEXT, mode TEXT, result_count INTEGER, created_at INTEGER, user_id TEXT, latency_ms REAL DEFAULT 0, top1_score REAL DEFAULT 0);
            CREATE TABLE IF NOT EXISTS request_metrics(
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              path TEXT NOT NULL,
              method TEXT NOT NULL,
              status_code INTEGER NOT NULL,
              latency_ms REAL NOT NULL,
              created_at INTEGER NOT NULL,
              user_id TEXT,
              error TEXT
            );
            CREATE TABLE IF NOT EXISTS interviews(id TEXT PRIMARY KEY, bank_id TEXT NOT NULL, user_id TEXT, question_ids TEXT NOT NULL, created_at INTEGER NOT NULL);
            CREATE TABLE IF NOT EXISTS interview_turns(id TEXT PRIMARY KEY, interview_id TEXT, question_id TEXT, prompt TEXT, answer TEXT, feedback TEXT, follow_up TEXT, depth INTEGER, score REAL, created_at INTEGER);
            CREATE TABLE IF NOT EXISTS push_settings(
              id INTEGER PRIMARY KEY CHECK(id=1),
              push_url TEXT,
              enabled INTEGER NOT NULL DEFAULT 1,
              push_time TEXT NOT NULL DEFAULT '09:00',
              bank_id TEXT,
              include_answer INTEGER NOT NULL DEFAULT 1,
              last_run_date TEXT,
              updated_at INTEGER
            );
            CREATE TABLE IF NOT EXISTS user_ai_settings(
              user_id TEXT PRIMARY KEY,
              provider TEXT NOT NULL DEFAULT 'deepseek',
              api_base TEXT NOT NULL,
              model TEXT NOT NULL,
              api_key TEXT NOT NULL,
              enabled INTEGER NOT NULL DEFAULT 1,
              created_at INTEGER NOT NULL,
              updated_at INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS user_push_settings(
              user_id TEXT PRIMARY KEY,
              push_url TEXT,
              enabled INTEGER NOT NULL DEFAULT 0,
              push_time TEXT,
              bank_id TEXT,
              include_answer INTEGER NOT NULL DEFAULT 1,
              last_run_date TEXT,
              updated_at INTEGER
            );
            CREATE TABLE IF NOT EXISTS coach_profiles(
              user_id TEXT PRIMARY KEY,
              target_position TEXT NOT NULL,
              interview_date TEXT,
              experience_level TEXT,
              daily_minutes INTEGER NOT NULL DEFAULT 30,
              jd_text TEXT,
              resume_summary TEXT,
              project_summary TEXT,
              focus_areas TEXT NOT NULL DEFAULT '[]',
              created_at INTEGER NOT NULL,
              updated_at INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS training_tasks(
              id TEXT PRIMARY KEY,
              user_id TEXT NOT NULL,
              task_type TEXT NOT NULL,
              title TEXT NOT NULL,
              description TEXT,
              due_date TEXT NOT NULL,
              status TEXT NOT NULL DEFAULT 'pending',
              action_url TEXT NOT NULL,
              source_ref TEXT NOT NULL DEFAULT '',
              created_at INTEGER NOT NULL,
              completed_at INTEGER,
              UNIQUE(user_id,due_date,task_type,title)
            );
            CREATE INDEX IF NOT EXISTS idx_training_tasks_user_due ON training_tasks(user_id,due_date,status);
            INSERT OR IGNORE INTO push_settings(id,enabled,push_time,include_answer,updated_at) VALUES(1,1,'09:00',1,0);
            """
        )
        admin_exists = connection.execute("SELECT 1 FROM users WHERE username=?", (ADMIN_USER,)).fetchone()
        if not admin_exists:
            connection.execute(
                "INSERT INTO users(username,password_hash,role,active,created_at) VALUES(?,?,?,?,?)",
                (ADMIN_USER, hash_password(ADMIN_PASSWORD), "admin", 1, int(time.time())),
            )
        migrate_owned_content(connection)
        add_column(connection, "search_logs", "latency_ms", "REAL DEFAULT 0")
        add_column(connection, "search_logs", "top1_score", "REAL DEFAULT 0")
        for column, definition in [
            ("bank_id", "TEXT"),
            ("tags", "TEXT DEFAULT '[]'"),
            ("chunk_index", "INTEGER DEFAULT 0"),
            ("parent_id", "TEXT"),
        ]:
            add_column(connection, "questions", column, definition)
        for column, definition in [
            ("stability", "REAL DEFAULT 1.0"),
            ("difficulty_factor", "REAL DEFAULT 5.0"),
            ("lapse_count", "INTEGER DEFAULT 0"),
            ("last_rating", "TEXT"),
        ]:
            add_column(connection, "user_question_states", column, definition)
        for column, definition in [
            ("mode", "TEXT DEFAULT 'general'"),
            ("project_title", "TEXT DEFAULT ''"),
            ("project_context", "TEXT DEFAULT ''"),
            ("generated_questions", "TEXT DEFAULT '[]'"),
            ("completed_at", "INTEGER"),
            ("report_score", "REAL"),
        ]:
            add_column(connection, "interviews", column, definition)
        add_column(connection, "search_logs", "user_id", "TEXT")
        # Keep metrics tables forward-compatible with databases created before
        # request attribution and error fields were introduced.
        add_column(connection, "request_metrics", "user_id", "TEXT")
        add_column(connection, "request_metrics", "error", "TEXT")
        default_id = ensure_default_bank(connection, ADMIN_USER)
        connection.execute(
            "UPDATE questions SET bank_id=?,owner_user_id=? WHERE bank_id IS NULL OR bank_id=''",
            (default_id, ADMIN_USER),
        )
        users = connection.execute("SELECT username FROM users").fetchall()
        for user in users:
            ensure_default_bank(connection, user["username"])
        connection.execute(
            """
            UPDATE user_push_settings SET bank_id=NULL
            WHERE bank_id IS NOT NULL AND NOT EXISTS(
              SELECT 1 FROM banks b WHERE b.id=user_push_settings.bank_id AND b.owner_user_id=user_push_settings.user_id
            )
            """
        )
        connection.execute(
            "UPDATE push_settings SET bank_id=NULL WHERE bank_id IS NOT NULL AND NOT EXISTS(SELECT 1 FROM banks b WHERE b.id=push_settings.bank_id AND b.owner_user_id=?)",
            (ADMIN_USER,),
        )
        connection.execute("UPDATE questions SET parent_id=NULL WHERE parent_id=id")


def create_token(user_id):
    token = secrets.token_urlsafe(32)
    with db() as connection:
        connection.execute(
            "INSERT INTO sessions(token,user_id,created_at,expires_at) VALUES(?,?,?,?)",
            (token, user_id, int(time.time()), int(time.time()) + 86400 * 7),
        )
    return token


def require_auth(request: Request, authorization: str | None = Header(default=None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "请先登录")
    token = authorization.removeprefix("Bearer ").strip()
    with db() as connection:
        row = connection.execute(
            """
            SELECT s.user_id,u.role,u.active
            FROM sessions s
            JOIN users u ON u.username=s.user_id
            WHERE s.token=? AND s.expires_at>?
            """,
            (token, int(time.time())),
        ).fetchone()
    if not row:
        raise HTTPException(401, "登录已过期")
    if not row["active"]:
        raise HTTPException(403, "账号已停用")
    request.state.user_id = row["user_id"]
    return row["user_id"]


def require_admin(user_id: str = Depends(require_auth)):
    with db() as connection:
        row = connection.execute("SELECT role FROM users WHERE username=? AND active=1", (user_id,)).fetchone()
    if not row or row["role"] != "admin":
        raise HTTPException(403, "仅管理员可以执行此操作")
    return user_id


def validate_ai_api_base(value):
    cleaned = str(value or "").strip().rstrip("/")
    parsed = urlparse(cleaned)
    localhost = parsed.hostname in {"localhost", "127.0.0.1", "::1"}
    if (
        parsed.scheme not in ({"http", "https"} if localhost else {"https"})
        or not parsed.hostname
        or parsed.username
        or parsed.password
        or parsed.params
        or parsed.query
        or parsed.fragment
    ):
        raise HTTPException(400, "API Base URL 必须是合法的 HTTPS 地址，且不能包含账号、查询参数或片段")
    return cleaned


def mask_api_key(value):
    value = str(value or "")
    if not value:
        return ""
    suffix = value[-4:] if len(value) > 4 else value[-1:]
    prefix = value[:3] if len(value) > 8 else "key"
    return f"{prefix}-****{suffix}"


def read_personal_ai_settings(user_id):
    with db() as connection:
        row = connection.execute("SELECT * FROM user_ai_settings WHERE user_id=?", (user_id,)).fetchone()
    return dict(row) if row else None


def user_role(user_id):
    with db() as connection:
        row = connection.execute("SELECT role FROM users WHERE username=?", (user_id,)).fetchone()
    return row["role"] if row else "user"


def resolve_ai_config(user_id):
    personal = read_personal_ai_settings(user_id)
    if personal:
        if not personal["enabled"]:
            return None
        return {
            "provider": personal["provider"],
            "api_base": personal["api_base"],
            "model": personal["model"],
            "api_key": personal["api_key"],
            "source": "personal",
        }
    if user_role(user_id) == "admin" and DEEPSEEK_API_KEY:
        return {
            "provider": "deepseek",
            "api_base": DEEPSEEK_BASE_URL.rstrip("/"),
            "model": DEEPSEEK_MODEL,
            "api_key": DEEPSEEK_API_KEY,
            "source": "server",
        }
    return None


def public_ai_settings(user_id):
    personal = read_personal_ai_settings(user_id)
    resolved = resolve_ai_config(user_id)
    source = resolved["source"] if resolved else "none"
    return {
        "configured": bool(resolved),
        "has_personal_config": bool(personal),
        "source": source,
        "provider": personal["provider"] if personal else (resolved or {}).get("provider", "deepseek"),
        "api_base": personal["api_base"] if personal else (resolved or {}).get("api_base", "https://api.deepseek.com"),
        "model": personal["model"] if personal else (resolved or {}).get("model", "deepseek-chat"),
        "api_key_masked": mask_api_key(personal["api_key"]) if personal else ("服务器托管" if source == "server" else ""),
        "enabled": bool(personal["enabled"]) if personal else bool(resolved),
        "updated_at": personal.get("updated_at") if personal else None,
        "credential_exposed": False,
    }


def read_file_push_url():
    if not SHOWDOC_PUSH_CONFIG.is_file():
        return None
    for raw_line in SHOWDOC_PUSH_CONFIG.read_text(encoding="utf-8").splitlines():
        key, separator, value = raw_line.strip().partition("=")
        if separator and key == "SHOWDOC_PUSH_URL":
            return value.strip().strip('"').strip("'")
    return None


def valid_showdoc_push_url(value):
    parsed = urlparse(str(value or "").strip())
    return (
        parsed.scheme == "https"
        and parsed.hostname == "push.showdoc.com.cn"
        and parsed.path.startswith("/server/api/push/")
        and len(parsed.path.removeprefix("/server/api/push/")) >= 20
        and not parsed.params
        and not parsed.query
        and not parsed.fragment
    )


def read_push_settings():
    with db() as connection:
        row = connection.execute("SELECT * FROM push_settings WHERE id=1").fetchone()
    return dict(row) if row else {"enabled": 1, "push_time": "09:00", "bank_id": None, "include_answer": 1, "last_run_date": None, "push_url": None}


def read_user_push_settings(user_id):
    with db() as connection:
        row = connection.execute("SELECT * FROM user_push_settings WHERE user_id=?", (user_id,)).fetchone()
    return dict(row) if row else None


def push_settings_for_user(user_id):
    personal = read_user_push_settings(user_id)
    if personal:
        return personal
    if user_role(user_id) == "admin":
        return read_push_settings()
    return {"enabled": 0, "push_time": None, "bank_id": None, "include_answer": 1, "last_run_date": None, "push_url": None}


def showdoc_push_url(user_id=None):
    settings = read_user_push_settings(user_id) if user_id else None
    if settings is None and user_id and user_role(user_id) != "admin":
        raise RuntimeError("当前账号尚未配置推送地址")
    settings = settings or read_push_settings()
    custom_url = str(settings.get("push_url") or "").strip()
    if custom_url:
        if valid_showdoc_push_url(custom_url):
            return custom_url
        raise RuntimeError("推送服务配置无效")
    if user_id and user_role(user_id) != "admin":
        raise RuntimeError("当前账号尚未配置推送地址")
    file_url = read_file_push_url()
    if file_url and valid_showdoc_push_url(file_url):
        return file_url
    raise RuntimeError("推送服务尚未配置")


def mask_push_url(url):
    if not url:
        return ""
    parsed = urlparse(url)
    token = parsed.path.rsplit("/", 1)[-1]
    return f"{parsed.scheme}://{parsed.netloc}/server/api/push/{'*' * 8}{token[-4:]}"


def next_push_at(settings):
    if not settings.get("enabled"):
        return None
    hour, minute = (int(item) for item in settings.get("push_time", "09:00").split(":", 1))
    now = datetime.now(PUSH_TIMEZONE)
    target = now.replace(hour=hour, minute=minute, second=0, microsecond=0)
    if target <= now:
        target += timedelta(days=1)
    return target.isoformat(timespec="minutes")


async def send_showdoc_push(title, content, user_id=None):
    try:
        url = showdoc_push_url(user_id)
    except RuntimeError as error:
        raise HTTPException(400, str(error)) from error
    try:
        async with httpx.AsyncClient(timeout=20) as client:
            response = await client.post(url, data={"title": title.strip(), "content": content.strip()})
        response.raise_for_status()
        result = response.json()
    except (OSError, httpx.HTTPError, ValueError) as error:
        raise HTTPException(502, "推送服务暂时不可用，请检查配置或稍后重试") from error
    if result.get("error_code") != 0:
        raise HTTPException(502, "推送服务拒绝了本次请求，请检查服务配置")
    return {"message": "消息已推送到微信"}


def clipped_push_text(value, limit):
    value = str(value or "").strip()
    return value if len(value) <= limit else value[:limit].rstrip() + "\n\n（内容过长，已截断）"


def choose_random_question(user_id, bank_id=None):
    where = ["q.parent_id IS NULL", "q.owner_user_id=?", "b.owner_user_id=?"]
    args = [user_id, user_id]
    if bank_id:
        where.append("b.id=?")
        args.append(bank_id)
    with db() as connection:
        rows = connection.execute(
            f"""
            SELECT q.id, q.question, q.answer, q.category, q.difficulty, q.position,
                   b.id AS bank_id, b.name AS bank_name
            FROM questions q
            JOIN banks b ON b.id=q.bank_id
            WHERE {' AND '.join(where)}
            """,
            args,
        ).fetchall()
    if not rows:
        return None
    return secrets.choice(rows)


def random_push_content(question, include_answer=True):
    title = f"鉴微随机题目｜{question['bank_name']}"
    lines = [
        f"# {question['bank_name']}",
        "",
        f"> 分类：{question['category'] or '未分类'} ｜ 难度：{question['difficulty'] or '未设置'} ｜ 岗位：{question['position'] or '通用'}",
        "",
        "## 题目",
        clipped_push_text(question["question"], 4000),
    ]
    if include_answer:
        lines.extend(["", "## 标准答案", clipped_push_text(question["answer"], 7000)])
    return title, "\n".join(lines)


async def push_random_question_message(bank_id=None, include_answer=True, user_id=None):
    effective_user = user_id or ADMIN_USER
    question = choose_random_question(effective_user, bank_id)
    if not question:
        raise RuntimeError("所选题库中没有可推送的题目" if bank_id else "当前没有可推送的题目")
    title, content = random_push_content(question, include_answer)
    await send_showdoc_push(title, content, user_id)
    return question


def normalize_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [item.strip() for item in re.split(r"[,，、\n]+", str(value)) if item.strip()]


def safe_json_list(value):
    try:
        parsed = json.loads(value or "[]")
        return parsed if isinstance(parsed, list) else []
    except Exception:
        return []


def searchable_text(item):
    get = item.get if isinstance(item, dict) else lambda key, default="": getattr(item, key, default)
    keywords = normalize_list(get("keywords", []))
    tags = normalize_list(get("tags", []))
    return "\n".join(
        [
            f"问题：{get('question')}",
            f"答案：{get('answer')}",
            f"分类：{get('category')}",
            f"岗位：{get('position')}",
            f"关键词：{' '.join(keywords)}",
            f"标签：{' '.join(tags)}",
        ]
    )


def hash_embedding(value):
    vector = [0.0] * DIMENSION
    text = re.sub(r"\s+", " ", value.lower())
    tokens = re.findall(r"[\w\u4e00-\u9fff]+", text)
    grams = tokens[:]
    for token in tokens:
        if len(token) > 2:
            grams.extend(token[index : index + 2] for index in range(len(token) - 1))
    for gram in grams:
        digest = hashlib.blake2b(gram.encode("utf-8"), digest_size=8).digest()
        index = int.from_bytes(digest[:4], "little") % DIMENSION
        sign = 1 if digest[4] % 2 == 0 else -1
        vector[index] += sign
    norm = math.sqrt(sum(item * item for item in vector)) or 1.0
    return [item / norm for item in vector]


def embed_document(value):
    if getattr(app.state, "embedder", None):
        return list(next(app.state.embedder.embed([value])))
    return hash_embedding(value)


def embed_query(value):
    if getattr(app.state, "embedder", None):
        return list(next(app.state.embedder.query_embed(value)))
    return hash_embedding(value)


def serialize(row):
    item = dict(row)
    item["keywords"] = safe_json_list(item.get("keywords"))
    item["tags"] = safe_json_list(item.get("tags"))
    item["children"] = item.get("children", [])
    return item


def split_answer(answer, size=900, overlap=120):
    answer = answer.strip()
    if len(answer) <= size:
        return [answer]
    chunks = []
    start = 0
    while start < len(answer):
        end = min(len(answer), start + size)
        chunk = answer[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(answer):
            break
        start = max(0, end - overlap)
    return chunks


def split_answer_for_import(answer, mode="smart", size=900, overlap=120):
    answer = str(answer or "").strip()
    size = max(300, min(int(size or 900), 4000))
    overlap = max(0, min(int(overlap or 0), min(size // 2, 600)))
    if mode == "none" or len(answer) <= size:
        return [answer]
    separators = [""] if mode == "fixed" else ["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""]
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=size,
        chunk_overlap=overlap,
        length_function=len,
        separators=separators,
        keep_separator=True,
    )
    return [chunk.strip() for chunk in splitter.split_text(answer) if chunk.strip()] or [answer]


def parse_markdown_import(content):
    normalized = content.replace("\r\n", "\n").replace("\r", "\n")
    current_bank = ""
    blocks = []
    current = []
    for line in normalized.splitlines():
        bank_match = re.match(r"^#\s+(.+?)\s*$", line)
        question_match = re.match(r"^##\s+(.+?)\s*$", line)
        separator_match = re.match(r"^\s*(?:---|\*\*\*)\s*$", line)
        if bank_match and not question_match:
            if current:
                blocks.append((current_bank, "\n".join(current)))
                current = []
            current_bank = bank_match.group(1).strip()
            continue
        if (question_match or separator_match) and current:
            blocks.append((current_bank, "\n".join(current)))
            current = []
        if separator_match:
            continue
        current.append(line)
    if current:
        blocks.append((current_bank, "\n".join(current)))

    field_map = {
        "\u5206\u7c7b": "category",
        "\u96be\u5ea6": "difficulty",
        "\u5c97\u4f4d": "position",
        "\u804c\u4f4d": "position",
        "\u5173\u952e\u8bcd": "keywords",
        "\u5173\u952e\u5b57": "keywords",
        "\u6807\u7b7e": "tags",
        "\u6765\u6e90": "source",
        "\u9898\u5e93": "bank_name",
        "\u9898\u5e93\u540d": "bank_name",
    }
    answer_headers = {"answer", "\u7b54\u6848", "\u6807\u51c6\u7b54\u6848", "\u53c2\u8003\u7b54\u6848"}

    def parse_meta_line(line):
        stripped = line.strip()
        if stripped.startswith(("-", "*")):
            stripped = stripped[1:].strip()
        separator = "\uff1a" if "\uff1a" in stripped else ":"
        if separator not in stripped:
            return None
        key, value = stripped.split(separator, 1)
        mapped = field_map.get(key.strip())
        if not mapped or not value.strip():
            return None
        return mapped, value.strip()

    items = []
    for bank_name, block in blocks:
        lines = block.strip().splitlines()
        if not lines:
            continue
        title = lines[0].strip()
        if title.startswith("##"):
            question = re.sub(r"^##\s+", "", title).strip()
            body_lines = lines[1:]
        else:
            question = title.lstrip("# ").strip()
            body_lines = lines[1:]
        if not question or not body_lines:
            continue

        answer_start = None
        for index, line in enumerate(body_lines):
            header = re.sub(r"^###\s+", "", line.strip()).strip().lower()
            if line.strip().startswith("###") and header in answer_headers:
                answer_start = index + 1
                break
        answer_lines = body_lines[answer_start:] if answer_start is not None else body_lines
        fields = {"question": question, "bank_name": bank_name}
        cleaned_answer_lines = []
        for line in answer_lines:
            parsed = parse_meta_line(line)
            if parsed:
                fields[parsed[0]] = parsed[1]
            else:
                cleaned_answer_lines.append(line)
        # Also allow metadata before answer header.
        for line in body_lines[: answer_start or 0]:
            parsed = parse_meta_line(line)
            if parsed:
                fields[parsed[0]] = parsed[1]
        answer = "\n".join(cleaned_answer_lines).strip()
        if not answer:
            continue
        fields["answer"] = answer
        if fields.get("bank_name"):
            fields["bank_name"] = str(fields["bank_name"]).strip()
        items.append(fields)
    return items


def parse_import(filename, raw):
    name = filename.lower()
    if name.endswith(".json"):
        data = json.loads(raw.decode("utf-8-sig"))
        return data if isinstance(data, list) else data.get("items", [])
    if name.endswith(".csv"):
        return list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    content = raw.decode("utf-8-sig")
    if name.endswith(".txt"):
        blocks = re.split(r"\n\s*\n", content)
        items = []
        for block in blocks:
            lines = [line.strip() for line in block.splitlines() if line.strip()]
            if len(lines) >= 2:
                items.append({"question": lines[0], "answer": "\n".join(lines[1:])})
        return items
    return parse_markdown_import(content)


def import_item_preview(raw_item, index, chunk_mode="smart", chunk_size=900, chunk_overlap=120):
    if not isinstance(raw_item, dict):
        return {"index": index, "valid": False, "error": "该记录不是对象", "question": f"第 {index} 条记录"}
    item = dict(raw_item)
    question = str(item.get("question") or item.get("题目") or "").strip()
    answer = str(item.get("answer") or item.get("答案") or item.get("标准答案") or "").strip()
    bank_name = str(item.get("bank_name") or item.get("bank") or item.get("题库") or "").strip()
    error = ""
    if len(question) < 2:
        error = "缺少有效题目"
    elif not answer:
        error = "缺少标准答案"
    chunks = split_answer_for_import(answer, chunk_mode, chunk_size, chunk_overlap) if answer else []
    return {
        "index": index,
        "valid": not error,
        "error": error,
        "question": question or f"第 {index} 条记录",
        "answer_preview": block_preview(answer, 180),
        "answer_length": len(answer),
        "category": str(item.get("category") or item.get("分类") or "未分类"),
        "difficulty": str(item.get("difficulty") or item.get("难度") or "中等"),
        "position": str(item.get("position") or item.get("岗位") or item.get("职位") or "通用"),
        "bank_name": bank_name,
        "estimated_chunks": len(chunks),
    }


def extract_project_document(filename, raw):
    name = (filename or "project.txt").lower()
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(raw))
            return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages if (page.extract_text() or "").strip())
        if name.endswith(".docx"):
            document = Document(io.BytesIO(raw))
            return "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
        if name.endswith(".json"):
            return json.dumps(json.loads(raw.decode("utf-8-sig")), ensure_ascii=False, indent=2)
        if name.endswith(".csv"):
            rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
            return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        return raw.decode("utf-8-sig")
    except Exception as error:
        raise HTTPException(400, f"Unable to parse project document: {error}") from error


def project_question_fallback(project_title):
    title = project_title or "\u8fd9\u4e2a\u9879\u76ee"
    prompts = [
        f"\u8bf7\u7528\u4e24\u5206\u949f\u8bf4\u6e05 {title} \u89e3\u51b3\u4e86\u4ec0\u4e48\u95ee\u9898\uff0c\u4f60\u7684\u8d21\u732e\u662f\u4ec0\u4e48\uff1f",
        f"{title} \u4e3a\u4ec0\u4e48\u9009\u62e9\u73b0\u5728\u7684\u6280\u672f\u65b9\u6848\uff0c\u66ff\u4ee3\u65b9\u6848\u662f\u4ec0\u4e48\uff1f",
        f"{title} \u6700\u96be\u7684\u4e00\u4e2a\u6280\u672f\u95ee\u9898\u662f\u4ec0\u4e48\uff0c\u4f60\u662f\u5982\u4f55\u5b9a\u4f4d\u548c\u89e3\u51b3\u7684\uff1f",
        f"\u5982\u679c {title} \u7684\u8bf7\u6c42\u91cf\u63d0\u5347\u5341\u500d\uff0c\u6700\u5148\u51fa\u95ee\u9898\u7684\u73af\u8282\u662f\u54ea\u91cc\uff1f",
        f"{title} \u7ebf\u4e0a\u51fa\u73b0\u8fc7\u4ec0\u4e48\u6545\u969c\uff0c\u5982\u4f55\u76d1\u63a7\u3001\u964d\u7ea7\u548c\u590d\u76d8\uff1f",
        f"\u73b0\u5728\u91cd\u505a {title}\uff0c\u4f60\u4f1a\u6539\u6389\u54ea\u4e09\u4e2a\u8bbe\u8ba1\uff0c\u4e3a\u4ec0\u4e48\uff1f",
    ]
    return [
        {
            "id": f"project-{uuid.uuid4()}",
            "question": prompt,
            "answer": "\u9700\u8981\u7ed3\u5408\u9879\u76ee\u4e8b\u5b9e\uff0c\u6309\u80cc\u666f\u3001\u9009\u578b\u3001\u6267\u884c\u3001\u7ed3\u679c\u548c\u590d\u76d8\u7ed3\u6784\u5316\u56de\u7b54\u3002",
            "category": "\u9879\u76ee\u6df1\u6316",
            "difficulty": "\u56f0\u96be",
            "position": "\u901a\u7528",
            "keywords": [],
            "tags": ["\u9879\u76ee\u6df1\u6316", "\u538b\u529b\u9762\u8bd5"],
            "source": "\u9879\u76ee\u6df1\u6316\u9762\u8bd5",
        }
        for prompt in prompts
    ]


def normalize_project_questions(raw_questions, project_title):
    normalized = []
    for raw in (raw_questions or [])[:6]:
        if not isinstance(raw, dict) or not str(raw.get("question", "")).strip():
            continue
        normalized.append(
            {
                "id": f"project-{uuid.uuid4()}",
                "question": str(raw["question"]).strip(),
                "answer": str(raw.get("reference_answer") or raw.get("answer") or "\u8bf7\u7ed3\u5408\u9879\u76ee\u4e8b\u5b9e\u56de\u7b54\u3002").strip(),
                "category": "\u9879\u76ee\u6df1\u6316",
                "difficulty": str(raw.get("difficulty") or "\u56f0\u96be"),
                "position": "\u901a\u7528",
                "keywords": normalize_list(raw.get("keywords", [])),
                "tags": ["\u9879\u76ee\u6df1\u6316", "\u538b\u529b\u9762\u8bd5", project_title][:3],
                "source": "\u9879\u76ee\u6df1\u6316\u9762\u8bd5",
                "evidence_target": str(raw.get("evidence_target") or "").strip(),
                "generated_by": "langchain-project-interview-agent",
            }
        )
    if len(normalized) < 6:
        normalized.extend(project_question_fallback(project_title)[len(normalized) : 6])
    return normalized[:6]


def insert_question_row(connection, payload, item_id, answer, chunk_index, parent_id):
    digest_source = payload["bank_id"] + "\n" + payload["question"].strip() + "\n" + answer.strip()
    if parent_id:
        digest_source += f"\n{parent_id}\n{chunk_index}"
    digest = hashlib.sha256(digest_source.encode()).hexdigest()
    now = int(time.time())
    connection.execute(
        """
        INSERT INTO questions(
            id,question,answer,category,difficulty,position,keywords,source,
            content_hash,created_at,updated_at,version,bank_id,tags,chunk_index,parent_id,owner_user_id
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            item_id,
            payload["question"].strip(),
            answer.strip(),
            payload["category"],
            payload["difficulty"],
            payload["position"],
            json.dumps(payload["keywords"], ensure_ascii=False),
            payload["source"],
            digest,
            now,
            now,
            1,
            payload["bank_id"],
            json.dumps(payload["tags"], ensure_ascii=False),
            chunk_index,
            parent_id,
            payload["owner_user_id"],
        ),
    )


def save_question(payload, user_id, chunk=True, chunk_mode="smart", chunk_size=900, chunk_overlap=120):
    if isinstance(payload, dict):
        payload = dict(payload)
        payload.setdefault("keywords", [])
        payload.setdefault("tags", [])
        payload.setdefault("category", "未分类")
        payload.setdefault("difficulty", "中等")
        payload.setdefault("position", "通用")
        payload.setdefault("source", "手动录入")
        payload["keywords"] = normalize_list(payload["keywords"])
        payload["tags"] = normalize_list(payload["tags"])
    payload = QuestionIn.model_validate(payload)
    bank_id = payload.bank_id or default_bank_id(user_id)
    require_owned_bank(bank_id, user_id)
    payload_dict = payload.model_dump()
    payload_dict["bank_id"] = bank_id
    payload_dict["owner_user_id"] = user_id
    root_id = str(uuid.uuid4())
    answer_chunks = split_answer_for_import(
        payload.answer,
        mode=chunk_mode if chunk else "none",
        size=chunk_size,
        overlap=chunk_overlap,
    )
    created_ids = [root_id]
    vector_rows = []
    child_rows = []
    if len(answer_chunks) == 1:
        vector_rows.append({"id": root_id, "owner_user_id": user_id, "embedding": embed_document(searchable_text({**payload_dict, "answer": payload.answer}))})
    else:
        for index, answer in enumerate(answer_chunks, 1):
            child_id = str(uuid.uuid4())
            child_rows.append((child_id, answer, index))
            created_ids.append(child_id)
            vector_rows.append({"id": child_id, "owner_user_id": user_id, "embedding": embed_document(searchable_text({**payload_dict, "answer": answer}))})
    with db() as connection:
        insert_question_row(connection, payload_dict, root_id, payload.answer, 0, None)
        for child_id, answer, index in child_rows:
            insert_question_row(connection, payload_dict, child_id, answer, index, root_id)
    try:
        app.state.milvus.insert(COLLECTION, vector_rows)
    except Exception:
        with db() as connection:
            placeholders = ",".join("?" for _ in created_ids)
            connection.execute(f"DELETE FROM questions WHERE id IN ({placeholders})", created_ids)
        raise
    return created_ids


def refresh_question_embedding(item_id):
    with db() as connection:
        row = connection.execute("SELECT * FROM questions WHERE id=?", (item_id,)).fetchone()
        children = connection.execute("SELECT id FROM questions WHERE parent_id=? ORDER BY chunk_index", (item_id,)).fetchall()
    if not row:
        return
    item = serialize(row)
    if item.get("parent_id"):
        vectors = [{"id": item_id, "owner_user_id": item["owner_user_id"], "embedding": embed_document(searchable_text(item))}]
        old_ids = [item_id]
    else:
        chunks = split_answer_for_import(item["answer"], "smart", 900, 120) if children else [item["answer"]]
        old_ids = [item_id, *[child["id"] for child in children]]
        child_rows = []
        vectors = []
        if len(chunks) == 1:
            vectors.append({"id": item_id, "owner_user_id": item["owner_user_id"], "embedding": embed_document(searchable_text(item))})
        else:
            for index, answer in enumerate(chunks, 1):
                child_id = str(uuid.uuid4())
                child_rows.append((child_id, answer, index))
                vectors.append({"id": child_id, "owner_user_id": item["owner_user_id"], "embedding": embed_document(searchable_text({**item, "answer": answer}))})
        with db() as connection:
            connection.execute("DELETE FROM questions WHERE parent_id=?", (item_id,))
            for child_id, answer, index in child_rows:
                insert_question_row(connection, item, child_id, answer, index, item_id)
    try:
        app.state.milvus.delete(COLLECTION, ids=old_ids)
    except Exception:
        pass
    app.state.milvus.insert(COLLECTION, vectors)


def key_score(query, item):
    query = query.lower().strip()
    body = searchable_text(item).lower()
    tokens = re.findall(r"[\w\u4e00-\u9fff]+", query) or list(query)
    if not tokens:
        return 0.0
    score = sum(token in body for token in tokens) / len(tokens)
    if query and query in body:
        score += 0.3
    return min(1.0, score)


def matches(item, payload):
    return (
        (not payload.bank_id or item.get("bank_id") == payload.bank_id)
        and (not payload.category or item["category"] == payload.category)
        and (not payload.difficulty or item["difficulty"] == payload.difficulty)
        and (not payload.position or item["position"] == payload.position)
    )


def ranked_search(payload, user_id):
    with db() as connection:
        rows = connection.execute("SELECT * FROM questions WHERE owner_user_id=?", (user_id,)).fetchall()
    items = {item["id"]: serialize(item) for item in rows}
    candidates = {}
    if payload.mode in ("semantic", "hybrid") and items:
        vector = embed_query(payload.query)
        try:
            hits = app.state.milvus.search(
                COLLECTION,
                data=[vector],
                anns_field="embedding",
                limit=min(max(payload.limit * 8, 30), 150),
                filter=f"owner_user_id == {json.dumps(user_id)}",
                search_params={"metric_type": "COSINE", "params": {}},
                output_fields=["id"],
            )[0]
        except Exception:
            app.state.milvus.load_collection(COLLECTION)
            hits = app.state.milvus.search(
                COLLECTION,
                data=[vector],
                anns_field="embedding",
                limit=min(max(payload.limit * 8, 30), 150),
                filter=f"owner_user_id == {json.dumps(user_id)}",
                search_params={"metric_type": "COSINE", "params": {}},
                output_fields=["id"],
            )[0]
        candidates = {item["id"]: max(0, min(1, float(item["distance"]))) for item in hits}
    if payload.mode == "keyword":
        candidates = {key: 0.0 for key in items}
    grouped = {}
    child_counts = {}
    for item in items.values():
        if item.get("parent_id"):
            child_counts[item["parent_id"]] = child_counts.get(item["parent_id"], 0) + 1
    total_weight = max(payload.semantic_weight + payload.keyword_weight, 0.01)
    for item_id, semantic in candidates.items():
        matched_item = items.get(item_id)
        if not matched_item:
            continue
        parent = items.get(matched_item.get("parent_id")) if matched_item.get("parent_id") else matched_item
        if not parent or not matches(parent, payload):
            continue
        keyword = key_score(payload.query, matched_item)
        if payload.mode == "keyword":
            score = keyword
        elif payload.mode == "semantic":
            score = semantic
        else:
            score = (semantic * payload.semantic_weight + keyword * payload.keyword_weight) / total_weight
        if score >= payload.min_score and (parent["id"] not in grouped or score > grouped[parent["id"]]["score"]):
            result = dict(parent)
            result["semantic_score"] = round(semantic, 4)
            result["keyword_score"] = round(keyword, 4)
            result["score"] = round(score, 4)
            result["chunk_count"] = child_counts.get(parent["id"], 0)
            if matched_item["id"] != parent["id"]:
                result["matched_chunk_id"] = matched_item["id"]
                result["matched_chunk"] = matched_item["answer"]
            grouped[parent["id"]] = result
    ranked = list(grouped.values())
    ranked.sort(key=lambda item: item["score"], reverse=True)
    return ranked[:payload.limit]


def create_vector_collection(client):
    schema = MilvusClient.create_schema(auto_id=False, enable_dynamic_field=False)
    schema.add_field("id", DataType.VARCHAR, is_primary=True, max_length=64)
    schema.add_field("owner_user_id", DataType.VARCHAR, max_length=128)
    schema.add_field("embedding", DataType.FLOAT_VECTOR, dim=DIMENSION)
    indexes = client.prepare_index_params()
    indexes.add_index("embedding", index_type="FLAT", metric_type="COSINE")
    client.create_collection(COLLECTION, schema=schema, index_params=indexes)


def rebuild_vector_index():
    """Recreate the active collection from SQLite; source data stays untouched."""
    with db() as connection:
        rows = connection.execute("SELECT * FROM questions ORDER BY created_at, id").fetchall()
    parent_ids_with_children = {row["parent_id"] for row in rows if row["parent_id"]}
    batch = []
    for row in rows:
        if row["id"] in parent_ids_with_children:
            continue
        item = serialize(row)
        batch.append({"id": item["id"], "owner_user_id": item["owner_user_id"], "embedding": embed_document(searchable_text(item))})
        if len(batch) >= 32:
            app.state.milvus.insert(COLLECTION, batch)
            batch = []
    if batch:
        app.state.milvus.insert(COLLECTION, batch)
    return len(rows)


def block_preview(text, limit=280):
    cleaned = re.sub(r"\s+", " ", text).strip()
    return cleaned[:limit] + ("…" if len(cleaned) > limit else "")


def chunk_lines(text):
    lines = [line.strip() for line in re.split(r"\n{2,}", text.strip()) if line.strip()]
    return lines or [text.strip()]


async def openai_compatible_json(system_prompt, user_prompt, ai_config):
    if not ai_config:
        raise HTTPException(400, "请先在 AI 配置页面设置并启用自己的 API")
    async with httpx.AsyncClient(timeout=120) as client:
        response = await client.post(
            f"{ai_config['api_base'].rstrip('/')}/chat/completions",
            headers={"Authorization": f"Bearer {ai_config['api_key']}", "Content-Type": "application/json"},
            json={
                "model": ai_config["model"],
                "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
                "response_format": {"type": "json_object"},
                "temperature": 0.25,
            },
        )
    if response.status_code >= 400:
        raise HTTPException(502, f"AI 服务调用失败（HTTP {response.status_code}），请检查个人配置")
    try:
        payload = response.json()["choices"][0]["message"]["content"]
        payload = re.sub(r"^```(?:json)?\s*|\s*```$", "", payload.strip(), flags=re.IGNORECASE)
        return json.loads(payload)
    except (KeyError, IndexError, TypeError, ValueError, json.JSONDecodeError) as error:
        raise HTTPException(502, "AI 服务返回格式无效，请检查模型是否支持 JSON 输出") from error


class ProjectInterviewAgent:
    def __init__(self, ai_config):
        if not ai_config:
            raise RuntimeError("AI API 未配置")
        self.model = ChatOpenAI(
            api_key=ai_config["api_key"],
            base_url=ai_config["api_base"],
            model=ai_config["model"],
            temperature=0.2,
            timeout=120,
            max_retries=2,
            model_kwargs={"response_format": {"type": "json_object"}},
        )
        self.parser = JsonOutputParser()

    async def generate_questions(self, title, context):
        analysis_prompt = ChatPromptTemplate.from_messages(
            [
                (
                    "system",
                    "你是资深技术面试 Agent 的项目分析阶段。只输出 JSON，不补充材料中不存在的事实。",
                ),
                (
                    "human",
                    "分析项目《{title}》。输出 JSON 字段 facts、candidate_contributions、architecture_decisions、incidents、metrics、risks、evidence_gaps，所有字段均为字符串数组。\n\n项目材料：\n{context}",
                ),
            ]
        )
        analysis = await (analysis_prompt | self.model | self.parser).ainvoke({"title": title, "context": context})
        question_prompt = ChatPromptTemplate.from_messages(
            [
                (
                    "system",
                    "你是严格的高级技术面试 Agent。根据项目事实规划递进式压力面试，只输出 JSON。问题必须验证候选人是否真实参与，覆盖个人贡献、架构取舍、故障定位、性能可靠性、安全和复盘，避免泛泛而谈。",
                ),
                (
                    "human",
                    "项目名称：{title}\n项目分析：{analysis}\n\n生成恰好 6 道递进问题。输出 questions 数组，每项包含 question、reference_answer、difficulty、keywords 数组和 evidence_target。标准答案只能依据分析结果，并明确期望候选人提供的数据或证据。",
                ),
            ]
        )
        generated = await (question_prompt | self.model | self.parser).ainvoke(
            {"title": title, "analysis": json.dumps(analysis, ensure_ascii=False)}
        )
        return generated, analysis

    async def evaluate_answer(self, interview, item, payload):
        prompt = ChatPromptTemplate.from_messages(
            [
                (
                    "system",
                    "你是严格但公平的项目深挖面试 Agent。根据项目事实验证回答真实性，识别含糊表述，并只输出 JSON。",
                ),
                (
                    "human",
                    """项目名称：{title}
项目材料：{context}
原始问题：{question}
参考要点：{reference}
本轮问题：{prompt}
候选人回答：{answer}
当前追问深度：{depth}

输出 JSON：score（0-100）、feedback、correct_answer、strengths 数组、weaknesses 数组、follow_up，以及 dimensions 对象（authenticity、architecture、troubleshooting、tradeoff、communication，均为 0-100）。follow_up 必须针对未证实的指标、个人贡献或取舍，最多两轮。""",
                ),
            ]
        )
        return await (prompt | self.model | self.parser).ainvoke(
            {
                "title": interview["project_title"],
                "context": str(interview["project_context"])[:16000],
                "question": item["question"],
                "reference": item["answer"],
                "prompt": payload.prompt,
                "answer": payload.answer,
                "depth": payload.depth,
            }
        )


def recall_probability(stability, last_reviewed_at):
    if not last_reviewed_at:
        return 0.0
    elapsed_days = max(0, (int(time.time()) - int(last_reviewed_at)) / 86400)
    return round(math.exp(-elapsed_days / max(float(stability or 1), 0.1)), 4)


def coach_today():
    return datetime.now(PUSH_TIMEZONE).date()


def serialize_coach_profile(row):
    if not row:
        return {
            "configured": False,
            "target_position": "",
            "interview_date": "",
            "experience_level": "1-3 年",
            "daily_minutes": 30,
            "jd_text": "",
            "resume_summary": "",
            "project_summary": "",
            "focus_areas": [],
        }
    profile = dict(row)
    profile["configured"] = True
    profile["focus_areas"] = safe_json_list(profile.get("focus_areas"))
    return profile


def refresh_training_plan(connection, user_id, profile):
    if not profile or not str(profile.get("target_position") or "").strip():
        return
    today = coach_today()
    focus = safe_json_list(profile.get("focus_areas"))
    focus_label = "、".join(str(item) for item in focus[:2]) or "岗位核心知识"
    has_project = bool(str(profile.get("project_summary") or "").strip())
    interview_url = "/?page=interview&mode=project" if has_project else "/?page=interview&mode=general"
    templates = [
        (0, "baseline", "完成一次基线模拟面试", "用完整的一轮回答建立当前能力基线。", interview_url),
        (0, "review", "清理今日到期复习", "优先巩固面试中已经暴露的薄弱题。", "/?page=review"),
        (1, "knowledge", f"专项梳理：{focus_label}", "通过有来源的 RAG 问答补齐核心概念和工程边界。", "/?page=ask"),
        (2, "project_interview", "完成一次项目深挖", "重点回答个人贡献、架构取舍、指标和故障复盘。", "/?page=interview&mode=project"),
        (3, "review", "完成薄弱题间隔复习", "不要背答案，先独立回答再核对评分要点。", "/?page=review"),
        (4, "general_interview", "完成一次岗位专项面试", "检查基础知识是否能在压力下完整表达。", "/?page=interview&mode=general"),
        (5, "knowledge", f"补齐薄弱知识：{focus_label}", "围绕最近低分项连续追问，形成可复述的答案。", "/?page=ask"),
        (6, "retest", "进行本周复测", "与基线结果比较，确认真正掌握而不是短期记忆。", interview_url),
    ]
    now = int(time.time())
    for offset, task_type, title, description, action_url in templates:
        due_date = (today + timedelta(days=offset)).isoformat()
        connection.execute(
            """
            INSERT OR IGNORE INTO training_tasks(
              id,user_id,task_type,title,description,due_date,status,action_url,source_ref,created_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?)
            """,
            (str(uuid.uuid4()), user_id, task_type, title, description, due_date, "pending", action_url, "coach-plan", now),
        )


def complete_next_interview_task(connection, user_id, mode, interview_id):
    today = coach_today().isoformat()
    types = ("project_interview", "baseline", "retest") if mode == "project" else ("general_interview", "baseline", "retest")
    placeholders = ",".join("?" for _ in types)
    row = connection.execute(
        f"""
        SELECT id FROM training_tasks
        WHERE user_id=? AND status='pending' AND due_date<=? AND task_type IN ({placeholders})
        ORDER BY due_date,created_at LIMIT 1
        """,
        (user_id, today, *types),
    ).fetchone()
    if row:
        connection.execute(
            "UPDATE training_tasks SET status='completed',completed_at=?,source_ref=? WHERE id=?",
            (int(time.time()), interview_id, row["id"]),
        )


def coach_dashboard(user_id):
    today = coach_today()
    now = int(time.time())
    with db() as connection:
        profile_row = connection.execute("SELECT * FROM coach_profiles WHERE user_id=?", (user_id,)).fetchone()
        profile = serialize_coach_profile(profile_row)
        knowledge_count = connection.execute(
            "SELECT COUNT(*) FROM questions WHERE owner_user_id=? AND parent_id IS NULL", (user_id,)
        ).fetchone()[0]
        interview_rows = connection.execute(
            """
            SELECT i.id,i.mode,i.project_title,i.created_at,i.completed_at,
              AVG(CASE WHEN t.depth=0 THEN t.score END) AS score,
              COUNT(CASE WHEN t.depth=0 THEN 1 END) AS answered
            FROM interviews i LEFT JOIN interview_turns t ON t.interview_id=i.id
            WHERE i.user_id=? GROUP BY i.id ORDER BY i.created_at DESC
            """,
            (user_id,),
        ).fetchall()
        review_stats = connection.execute(
            """
            SELECT COUNT(*) AS learned,
              SUM(CASE WHEN next_review_at IS NULL OR next_review_at<=? THEN 1 ELSE 0 END) AS due,
              COALESCE(SUM(review_count),0) AS completed
            FROM user_question_states WHERE user_id=?
            """,
            (now, user_id),
        ).fetchone()
        tasks = connection.execute(
            """
            SELECT * FROM training_tasks
            WHERE user_id=? AND (due_date<=? OR (status='completed' AND due_date=?))
            ORDER BY CASE status WHEN 'pending' THEN 0 ELSE 1 END,due_date,created_at LIMIT 12
            """,
            (user_id, today.isoformat(), today.isoformat()),
        ).fetchall()
        completed_task_count = connection.execute(
            "SELECT COUNT(*) FROM training_tasks WHERE user_id=? AND status='completed'", (user_id,)
        ).fetchone()[0]
        feedback_rows = connection.execute(
            """
            SELECT t.feedback FROM interview_turns t JOIN interviews i ON i.id=t.interview_id
            WHERE i.user_id=? AND t.depth=0 ORDER BY t.created_at DESC LIMIT 60
            """,
            (user_id,),
        ).fetchall()

    completed_interviews = [row for row in interview_rows if int(row["answered"] or 0) > 0]
    scores = [float(row["score"] or 0) for row in completed_interviews]
    average_score = round(sum(scores) / len(scores), 1) if scores else 0
    dimension_values = {}
    weakness_counts = {}
    for row in feedback_rows:
        try:
            evaluation = json.loads(row["feedback"] or "{}")
        except (TypeError, json.JSONDecodeError):
            continue
        for key, value in (evaluation.get("dimensions") or {}).items():
            try:
                dimension_values.setdefault(key, []).append(float(value))
            except (TypeError, ValueError):
                continue
        for item in evaluation.get("weaknesses", []):
            label = str(item).strip()
            if label:
                weakness_counts[label] = weakness_counts.get(label, 0) + 1

    dimension_labels = {
        "authenticity": "项目真实性与证据",
        "architecture": "架构设计",
        "troubleshooting": "故障排查",
        "tradeoff": "技术取舍",
        "communication": "表达结构",
    }
    dimensions = {
        key: round(sum(values) / len(values), 1) for key, values in dimension_values.items() if values
    }
    weak_areas = [
        {"name": dimension_labels.get(key, key), "score": score, "source": "能力维度"}
        for key, score in sorted(dimensions.items(), key=lambda item: item[1])[:3]
    ]
    for label, count in sorted(weakness_counts.items(), key=lambda item: (-item[1], item[0])):
        if len(weak_areas) >= 5:
            break
        if not any(item["name"] == label for item in weak_areas):
            weak_areas.append({"name": label, "score": None, "source": f"出现 {count} 次"})

    profile_score = 0
    if profile["configured"]:
        profile_score = 8
        profile_score += 4 if profile.get("interview_date") else 0
        profile_score += 4 if profile.get("jd_text") else 0
        profile_score += 4 if profile.get("resume_summary") or profile.get("project_summary") else 0
    knowledge_score = min(20, round(knowledge_count / 30 * 20))
    interview_score = min(35, round(len(completed_interviews) * 5 + average_score * 0.2))
    review_score = min(25, round(float(review_stats["completed"] or 0) * 2.5))
    readiness = min(100, profile_score + knowledge_score + interview_score + review_score)
    days_left = None
    if profile.get("interview_date"):
        try:
            days_left = max(0, (datetime.strptime(profile["interview_date"], "%Y-%m-%d").date() - today).days)
        except ValueError:
            pass

    recent = []
    for row in reversed(completed_interviews[:8]):
        recent.append(
            {
                "date": datetime.fromtimestamp(row["created_at"], PUSH_TIMEZONE).strftime("%m-%d"),
                "score": round(float(row["score"] or 0), 1),
                "mode": row["mode"] or "general",
            }
        )
    return {
        "profile": profile,
        "readiness": readiness,
        "readiness_breakdown": {
            "目标材料": profile_score,
            "知识储备": knowledge_score,
            "模拟面试": interview_score,
            "复习巩固": review_score,
        },
        "days_left": days_left,
        "stats": {
            "knowledge_count": knowledge_count,
            "interview_count": len(completed_interviews),
            "average_score": average_score,
            "due_reviews": int(review_stats["due"] or 0),
            "completed_tasks": completed_task_count,
        },
        "tasks": [dict(row) for row in tasks],
        "weak_areas": weak_areas,
        "dimensions": dimensions,
        "recent_interviews": recent,
    }


def get_question_tree(connection, bank_id=None):
    where = "WHERE bank_id=?" if bank_id else ""
    args = [bank_id] if bank_id else []
    rows = connection.execute(
        f"""
        SELECT parent_id, COUNT(*) AS count
        FROM questions
        {where}
        GROUP BY parent_id
        """,
        args,
    ).fetchall()
    return {row["parent_id"] or "": row["count"] for row in rows}


async def push_scheduler_loop():
    while True:
        try:
            settings = read_push_settings()
            now = datetime.now(PUSH_TIMEZONE)
            today = now.date().isoformat()
            if settings.get("enabled") and settings.get("push_time") == now.strftime("%H:%M") and settings.get("last_run_date") != today:
                with db() as connection:
                    changed = connection.execute(
                        "UPDATE push_settings SET last_run_date=?, updated_at=? WHERE id=1 AND last_run_date IS NOT ?",
                        (today, int(time.time()), today),
                    ).rowcount
                if changed:
                    try:
                        await push_random_question_message(settings.get("bank_id"), bool(settings.get("include_answer", 1)))
                        logger.info("scheduled random question push completed")
                    except Exception:
                        with db() as connection:
                            connection.execute("UPDATE push_settings SET last_run_date=NULL WHERE id=1 AND last_run_date=?", (today,))
                        logger.exception("scheduled random question push failed")
            # Personal schedules are isolated per ordinary user and never use
            # the server-level ShowDoc credential.
            with db() as connection:
                personal_rows = connection.execute(
                    """
                    SELECT s.*,u.role FROM user_push_settings s
                    JOIN users u ON u.username=s.user_id
                    WHERE s.enabled=1 AND u.active=1 AND u.role!='admin'
                    """
                ).fetchall()
            for personal in personal_rows:
                personal = dict(personal)
                if personal.get("push_time") != now.strftime("%H:%M") or personal.get("last_run_date") == today:
                    continue
                with db() as connection:
                    changed = connection.execute(
                        "UPDATE user_push_settings SET last_run_date=?, updated_at=? WHERE user_id=? AND last_run_date IS NOT ?",
                        (today, int(time.time()), personal["user_id"], today),
                    ).rowcount
                if changed:
                    try:
                        await push_random_question_message(
                            personal.get("bank_id"), bool(personal.get("include_answer", 1)), personal["user_id"]
                        )
                        logger.info("personal scheduled push completed for user")
                    except Exception:
                        with db() as connection:
                            connection.execute(
                                "UPDATE user_push_settings SET last_run_date=NULL WHERE user_id=? AND last_run_date=?",
                                (personal["user_id"], today),
                            )
                        logger.exception("personal scheduled push failed")
        except Exception:
            logger.exception("push scheduler loop failed")
        await asyncio.sleep(20)


@asynccontextmanager
async def lifespan(app):
    init_db()
    app.state.embedder = None if EMBEDDING_BACKEND == "hash" else TextEmbedding(model_name=MODEL, cache_dir=FASTEMBED_CACHE_PATH)
    app.state.milvus = MilvusClient(uri=str(VDB))
    if REBUILD_VECTOR_INDEX and app.state.milvus.has_collection(COLLECTION):
        app.state.milvus.drop_collection(COLLECTION)
    created = not app.state.milvus.has_collection(COLLECTION)
    if created:
        create_vector_collection(app.state.milvus)
    app.state.milvus.load_collection(COLLECTION)
    if created:
        rebuild_vector_index()
    app.state.push_scheduler = asyncio.create_task(push_scheduler_loop())
    try:
        yield
    finally:
        app.state.push_scheduler.cancel()
        with suppress(asyncio.CancelledError):
            await app.state.push_scheduler


app = FastAPI(title="鉴微", version=APP_VERSION, lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


@app.middleware("http")
async def collect_request_metrics(request: Request, call_next):
    started = time.perf_counter()
    status_code = 500
    error = None
    try:
        response = await call_next(request)
        status_code = response.status_code
        return response
    except Exception as exc:
        error = type(exc).__name__
        raise
    finally:
        path = request.url.path
        if path.startswith("/api/"):
            try:
                with db() as connection:
                    connection.execute(
                        "INSERT INTO request_metrics(path,method,status_code,latency_ms,created_at,user_id,error) VALUES(?,?,?,?,?,?,?)",
                        (path, request.method, status_code, round((time.perf_counter() - started) * 1000, 2), int(time.time()), getattr(request.state, "user_id", None), error),
                    )
            except Exception:
                logger.exception("failed to persist request metric")


@app.post("/api/auth/login")
def login(payload: LoginIn):
    username = payload.username.strip()
    with db() as connection:
        user = connection.execute("SELECT * FROM users WHERE LOWER(username)=LOWER(?)", (username,)).fetchone()
    if not user or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(401, "用户名或密码错误")
    if not user["active"]:
        raise HTTPException(403, "账号已停用，请联系管理员")
    with db() as connection:
        connection.execute("UPDATE users SET last_login_at=? WHERE username=?", (int(time.time()), user["username"]))
    return {"token": create_token(user["username"]), "user_id": user["username"], "role": user["role"]}


@app.post("/api/auth/register", status_code=201)
def register(payload: RegisterIn):
    username = payload.username.strip()
    code_hash = invite_code_hash(payload.invite_code)
    now = int(time.time())
    with db() as connection:
        connection.execute("BEGIN IMMEDIATE")
        exists = connection.execute("SELECT 1 FROM users WHERE LOWER(username)=LOWER(?)", (username,)).fetchone()
        if exists:
            raise HTTPException(409, "用户名已存在")
        invite = connection.execute("SELECT * FROM invites WHERE code_hash=?", (code_hash,)).fetchone()
        if not invite or invite["revoked_at"] or invite["used_at"] or invite["expires_at"] <= now:
            raise HTTPException(400, "邀请码无效、已使用或已过期")
        connection.execute(
            "INSERT INTO users(username,password_hash,role,active,invited_by,created_at) VALUES(?,?,?,?,?,?)",
            (username, hash_password(payload.password), "user", 1, invite["created_by"], now),
        )
        connection.execute(
            "UPDATE invites SET used_by=?,used_at=? WHERE id=?",
            (username, now, invite["id"]),
        )
        ensure_default_bank(connection, username)
    return {"token": create_token(username), "user_id": username, "role": "user", "message": "注册成功"}


@app.get("/api/auth/me")
def auth_me(user_id: str = Depends(require_auth)):
    with db() as connection:
        user = connection.execute(
            "SELECT username,role,active,created_at,last_login_at FROM users WHERE username=?",
            (user_id,),
        ).fetchone()
    return dict(user)


@app.get("/api/coach/dashboard")
def get_coach_dashboard(user_id: str = Depends(require_auth)):
    return coach_dashboard(user_id)


@app.get("/api/coach/profile")
def get_coach_profile(user_id: str = Depends(require_auth)):
    with db() as connection:
        row = connection.execute("SELECT * FROM coach_profiles WHERE user_id=?", (user_id,)).fetchone()
    return serialize_coach_profile(row)


@app.put("/api/coach/profile")
def save_coach_profile(payload: CoachProfileIn, user_id: str = Depends(require_auth)):
    interview_date = payload.interview_date.strip()
    if interview_date:
        try:
            parsed_date = datetime.strptime(interview_date, "%Y-%m-%d").date()
        except ValueError as error:
            raise HTTPException(400, "面试日期格式应为 YYYY-MM-DD") from error
        if parsed_date < coach_today():
            raise HTTPException(400, "面试日期不能早于今天")
    now = int(time.time())
    focus_areas = list(dict.fromkeys(item.strip() for item in payload.focus_areas if item.strip()))[:12]
    with db() as connection:
        connection.execute(
            """
            INSERT INTO coach_profiles(
              user_id,target_position,interview_date,experience_level,daily_minutes,
              jd_text,resume_summary,project_summary,focus_areas,created_at,updated_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(user_id) DO UPDATE SET
              target_position=excluded.target_position,interview_date=excluded.interview_date,
              experience_level=excluded.experience_level,daily_minutes=excluded.daily_minutes,
              jd_text=excluded.jd_text,resume_summary=excluded.resume_summary,
              project_summary=excluded.project_summary,focus_areas=excluded.focus_areas,
              updated_at=excluded.updated_at
            """,
            (
                user_id,
                payload.target_position.strip(),
                interview_date,
                payload.experience_level.strip(),
                payload.daily_minutes,
                payload.jd_text.strip(),
                payload.resume_summary.strip(),
                payload.project_summary.strip(),
                json.dumps(focus_areas, ensure_ascii=False),
                now,
                now,
            ),
        )
        row = connection.execute("SELECT * FROM coach_profiles WHERE user_id=?", (user_id,)).fetchone()
        connection.execute(
            "DELETE FROM training_tasks WHERE user_id=? AND status='pending' AND source_ref='coach-plan' AND due_date>=?",
            (user_id, coach_today().isoformat()),
        )
        refresh_training_plan(connection, user_id, dict(row))
    return {"message": "求职目标与训练计划已更新", "profile": serialize_coach_profile(row)}


@app.post("/api/coach/materials/extract")
async def extract_coach_material(
    file: UploadFile = File(...),
    user_id: str = Depends(require_auth),
):
    if not file.filename or not file.filename.lower().endswith((".pdf", ".docx", ".md", ".markdown", ".txt", ".json", ".csv")):
        raise HTTPException(400, "仅支持 PDF、Word、Markdown、TXT、JSON 和 CSV 文件")
    raw = await file.read()
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(400, "材料文件不能超过 10 MB")
    text = extract_project_document(file.filename, raw).strip()
    if not text:
        raise HTTPException(400, "未能从材料中提取有效文字")
    return {"filename": file.filename, "text": text[:30000], "characters": min(len(text), 30000)}


@app.patch("/api/coach/tasks/{task_id}")
def update_coach_task(task_id: str, payload: CoachTaskIn, user_id: str = Depends(require_auth)):
    status = "completed" if payload.completed else "pending"
    completed_at = int(time.time()) if payload.completed else None
    with db() as connection:
        changed = connection.execute(
            "UPDATE training_tasks SET status=?,completed_at=? WHERE id=? AND user_id=?",
            (status, completed_at, task_id, user_id),
        ).rowcount
    if not changed:
        raise HTTPException(404, "训练任务不存在")
    return {"message": "训练进度已更新", "status": status}


@app.get("/api/ai/settings")
def get_ai_settings(user_id: str = Depends(require_auth)):
    return public_ai_settings(user_id)


@app.put("/api/ai/settings")
def update_ai_settings(payload: UserAISettingsIn, user_id: str = Depends(require_auth)):
    api_base = validate_ai_api_base(payload.api_base)
    model = payload.model.strip()
    supplied_key = str(payload.api_key or "").strip()
    if not model:
        raise HTTPException(400, "模型名称不能为空")
    with db() as connection:
        existing = connection.execute("SELECT api_key,created_at FROM user_ai_settings WHERE user_id=?", (user_id,)).fetchone()
        api_key = supplied_key or (existing["api_key"] if existing else "")
        if not api_key:
            raise HTTPException(400, "首次配置必须填写 API Key")
        now = int(time.time())
        connection.execute(
            """
            INSERT INTO user_ai_settings(user_id,provider,api_base,model,api_key,enabled,created_at,updated_at)
            VALUES(?,?,?,?,?,?,?,?)
            ON CONFLICT(user_id) DO UPDATE SET
              provider=excluded.provider,
              api_base=excluded.api_base,
              model=excluded.model,
              api_key=excluded.api_key,
              enabled=excluded.enabled,
              updated_at=excluded.updated_at
            """,
            (user_id, payload.provider, api_base, model, api_key, int(payload.enabled), existing["created_at"] if existing else now, now),
        )
    return {**public_ai_settings(user_id), "message": "个人 AI 配置已保存", "persisted": True}


@app.delete("/api/ai/settings")
def delete_ai_settings(user_id: str = Depends(require_auth)):
    with db() as connection:
        connection.execute("DELETE FROM user_ai_settings WHERE user_id=?", (user_id,))
    return {**public_ai_settings(user_id), "message": "个人 AI 配置已删除"}


@app.post("/api/ai/settings/test")
async def test_ai_settings(payload: UserAISettingsIn, user_id: str = Depends(require_auth)):
    api_base = validate_ai_api_base(payload.api_base)
    existing = read_personal_ai_settings(user_id)
    api_key = str(payload.api_key or "").strip() or (existing or {}).get("api_key", "")
    if not api_key:
        raise HTTPException(400, "请填写 API Key 后再测试")
    config = {
        "provider": payload.provider,
        "api_base": api_base,
        "model": payload.model.strip(),
        "api_key": api_key,
        "source": "personal",
    }
    result = await openai_compatible_json(
        "你是连接测试助手，只输出 JSON。",
        '只输出 {"status":"ok"}，不要添加其他内容。',
        config,
    )
    if result.get("status") != "ok":
        raise HTTPException(502, "AI 服务已响应，但返回内容不符合连接测试要求")
    return {"ok": True, "message": "连接测试成功", "model": config["model"]}


@app.get("/api/admin/invites")
def admin_invites(_: str = Depends(require_admin)):
    now = int(time.time())
    with db() as connection:
        rows = connection.execute("SELECT * FROM invites ORDER BY created_at DESC").fetchall()
    items = []
    for row in rows:
        item = dict(row)
        if item["revoked_at"]:
            status = "revoked"
        elif item["used_at"]:
            status = "used"
        elif item["expires_at"] <= now:
            status = "expired"
        else:
            status = "active"
        items.append(
            {
                "id": item["id"],
                "code_masked": item["code_prefix"] + "-*****-*****-*****",
                "note": item["note"] or "",
                "created_by": item["created_by"],
                "created_at": item["created_at"],
                "expires_at": item["expires_at"],
                "used_by": item["used_by"],
                "used_at": item["used_at"],
                "status": status,
            }
        )
    return {"items": items, "total": len(items)}


@app.post("/api/admin/invites", status_code=201)
def create_invite(payload: InviteCreateIn, admin_id: str = Depends(require_admin)):
    code = new_invite_code()
    now = int(time.time())
    item_id = str(uuid.uuid4())
    with db() as connection:
        connection.execute(
            "INSERT INTO invites(id,code_hash,code_prefix,note,created_by,created_at,expires_at) VALUES(?,?,?,?,?,?,?)",
            (item_id, invite_code_hash(code), code[:8], payload.note.strip(), admin_id, now, now + payload.expires_in_days * 86400),
        )
    return {
        "id": item_id,
        "code": code,
        "note": payload.note.strip(),
        "expires_at": now + payload.expires_in_days * 86400,
        "message": "邀请码已创建，请立即保存",
    }


@app.delete("/api/admin/invites/{invite_id}")
def revoke_invite(invite_id: str, _: str = Depends(require_admin)):
    with db() as connection:
        updated = connection.execute(
            "UPDATE invites SET revoked_at=? WHERE id=? AND used_at IS NULL AND revoked_at IS NULL",
            (int(time.time()), invite_id),
        ).rowcount
    if not updated:
        raise HTTPException(400, "邀请码不存在、已使用或已撤销")
    return {"message": "邀请码已撤销"}


@app.get("/api/admin/users")
def admin_users(_: str = Depends(require_admin)):
    with db() as connection:
        rows = connection.execute(
            "SELECT username,role,active,invited_by,created_at,last_login_at FROM users ORDER BY created_at DESC"
        ).fetchall()
    return {"items": [dict(row) for row in rows], "total": len(rows)}


@app.patch("/api/admin/users/{username}")
def update_user_status(username: str, payload: UserStatusIn, _: str = Depends(require_admin)):
    with db() as connection:
        user = connection.execute("SELECT role FROM users WHERE username=?", (username,)).fetchone()
        if not user:
            raise HTTPException(404, "用户不存在")
        if user["role"] == "admin":
            raise HTTPException(400, "不能停用管理员账号")
        connection.execute("UPDATE users SET active=? WHERE username=?", (int(payload.active), username))
        if not payload.active:
            connection.execute("DELETE FROM sessions WHERE user_id=?", (username,))
    return {"message": "用户状态已更新", "active": payload.active}


def percentile(values, ratio):
    values = sorted(float(value) for value in values if value is not None)
    if not values:
        return 0
    index = min(len(values) - 1, max(0, math.ceil(len(values) * ratio) - 1))
    return round(values[index], 2)


@app.get("/api/admin/metrics")
def admin_metrics(_: str = Depends(require_admin)):
    now = int(time.time())
    windows = {"24h": now - 86400, "7d": now - 7 * 86400}
    result = {"generated_at": now, "windows": {}}
    with db() as connection:
        knowledge = connection.execute(
            "SELECT COUNT(*) total, SUM(CASE WHEN parent_id IS NOT NULL THEN 1 ELSE 0 END) chunks FROM questions"
        ).fetchone()
        banks = connection.execute("SELECT COUNT(*) FROM banks").fetchone()[0]
        users = connection.execute("SELECT COUNT(*) FROM users WHERE active=1").fetchone()[0]
        for label, since in windows.items():
            requests = connection.execute(
                "SELECT status_code,latency_ms FROM request_metrics WHERE created_at>=? ORDER BY latency_ms",
                (since,),
            ).fetchall()
            searches = connection.execute(
                "SELECT mode,result_count,latency_ms,top1_score FROM search_logs WHERE created_at>=?",
                (since,),
            ).fetchall()
            total = len(requests)
            errors = sum(1 for row in requests if int(row["status_code"]) >= 500)
            latencies = [row["latency_ms"] for row in requests]
            search_count = len(searches)
            empty = sum(1 for row in searches if int(row["result_count"] or 0) == 0)
            result["windows"][label] = {
                "requests": total,
                "errors_5xx": errors,
                "error_rate": round(errors / total, 4) if total else 0,
                "latency_ms": {
                    "avg": round(sum(latencies) / total, 2) if total else 0,
                    "p50": percentile(latencies, 0.50),
                    "p95": percentile(latencies, 0.95),
                    "p99": percentile(latencies, 0.99),
                },
                "searches": search_count,
                "empty_recall_rate": round(empty / search_count, 4) if search_count else 0,
                "avg_top1_score": round(sum(float(row["top1_score"] or 0) for row in searches) / search_count, 4) if search_count else 0,
                "avg_search_latency_ms": round(sum(float(row["latency_ms"] or 0) for row in searches) / search_count, 2) if search_count else 0,
                "modes": {
                    mode: sum(1 for row in searches if row["mode"] == mode)
                    for mode in sorted({row["mode"] for row in searches})
                },
            }
    result["knowledge"] = {"questions": int(knowledge["total"] or 0), "chunks": int(knowledge["chunks"] or 0), "banks": int(banks)}
    result["users"] = {"active": int(users)}
    result["targets"] = {
        "availability": ">=99.5%",
        "api_p95_ms": "<=5000",
        "empty_recall_rate": "<5%",
        "top1_score": ">=80%",
    }
    return result


@app.get("/health")
def health():
    return {
        "status": "ok",
        "name": "鉴微",
        "version": APP_VERSION,
        "model": MODEL,
        "embedding_backend": EMBEDDING_BACKEND,
        "vector_collection": COLLECTION,
        "llm": DEEPSEEK_MODEL,
    }


@app.get("/api/banks")
def banks(user_id: str = Depends(require_auth)):
    with db() as connection:
        rows = connection.execute(
            "SELECT b.*,COUNT(q.id) question_count FROM banks b LEFT JOIN questions q ON q.bank_id=b.id AND q.parent_id IS NULL AND q.owner_user_id=? WHERE b.owner_user_id=? GROUP BY b.id ORDER BY b.created_at",
            (user_id, user_id),
        ).fetchall()
    return {"items": [dict(row) for row in rows]}


@app.get("/api/push/status")
def push_status(user_id: str = Depends(require_auth)):
    settings = push_settings_for_user(user_id)
    try:
        configured_url = showdoc_push_url(user_id)
        configured = True
    except RuntimeError:
        configured_url = ""
        configured = False
    return {
        "configured": configured,
        "provider": "ShowDoc",
        "credential_exposed": False,
        "push_url_masked": mask_push_url(configured_url),
        "enabled": bool(settings.get("enabled")),
        "push_time": settings.get("push_time", "09:00"),
        "next_push_at": next_push_at(settings),
    }


@app.get("/api/push/settings")
def get_push_settings(user_id: str = Depends(require_auth)):
    settings = push_settings_for_user(user_id)
    try:
        configured_url = showdoc_push_url(user_id)
    except RuntimeError:
        configured_url = ""
    return {
        "configured": bool(configured_url),
        "provider": "ShowDoc",
        "credential_exposed": False,
        "push_url_masked": mask_push_url(configured_url),
        "enabled": bool(settings.get("enabled")),
        "push_time": settings.get("push_time"),
        "bank_id": settings.get("bank_id"),
        "include_answer": bool(settings.get("include_answer", 1)),
        "last_run_date": settings.get("last_run_date"),
        "next_push_at": next_push_at(settings),
    }


@app.put("/api/push/settings")
def update_push_settings(payload: PushSettingsIn, user_id: str = Depends(require_auth)):
    push_url = None
    if payload.push_url is not None and payload.push_url.strip():
        push_url = payload.push_url.strip()
        if not valid_showdoc_push_url(push_url):
            raise HTTPException(400, "推送网址必须是合法的 ShowDoc HTTPS 推送地址")
    if payload.bank_id:
        require_owned_bank(payload.bank_id, user_id)
    now = int(time.time())
    if user_role(user_id) == "admin":
        with db() as connection:
            if push_url is None:
                connection.execute(
                    "UPDATE push_settings SET enabled=?, push_time=?, bank_id=?, include_answer=?, updated_at=? WHERE id=1",
                    (int(payload.enabled), payload.push_time, payload.bank_id, int(payload.include_answer), now),
                )
            else:
                connection.execute(
                    "UPDATE push_settings SET push_url=?, enabled=?, push_time=?, bank_id=?, include_answer=?, updated_at=? WHERE id=1",
                    (push_url, int(payload.enabled), payload.push_time, payload.bank_id, int(payload.include_answer), now),
                )
    else:
        with db() as connection:
            existing = connection.execute("SELECT push_url FROM user_push_settings WHERE user_id=?", (user_id,)).fetchone()
            target_url = push_url if push_url is not None else (existing["push_url"] if existing else None)
            connection.execute(
                """
                INSERT INTO user_push_settings(user_id,push_url,enabled,push_time,bank_id,include_answer,updated_at)
                VALUES(?,?,?,?,?,?,?)
                ON CONFLICT(user_id) DO UPDATE SET
                  push_url=excluded.push_url, enabled=excluded.enabled,
                  push_time=excluded.push_time, bank_id=excluded.bank_id,
                  include_answer=excluded.include_answer, updated_at=excluded.updated_at
                """,
                (user_id, target_url, int(payload.enabled), payload.push_time, payload.bank_id, int(payload.include_answer), now),
            )
    return get_push_settings(user_id)


@app.post("/api/push/custom")
async def push_custom(payload: CustomPushIn, user_id: str = Depends(require_auth)):
    return await send_showdoc_push(payload.title, payload.content, user_id)


@app.post("/api/push/random-question")
async def push_random_question(payload: RandomQuestionPushIn, user_id: str = Depends(require_auth)):
    try:
        question = await push_random_question_message(payload.bank_id, payload.include_answer, user_id)
    except RuntimeError as error:
        raise HTTPException(404, str(error)) from error
    result = {"message": "消息已推送到微信"}
    result.update(
        {
            "bank": {"id": question["bank_id"], "name": question["bank_name"]},
            "question": {"id": question["id"], "title": question["question"]},
            "included_answer": payload.include_answer,
        }
    )
    return result


@app.post("/api/banks", status_code=201)
def create_bank(payload: BankIn, user_id: str = Depends(require_auth)):
    bank_id = str(uuid.uuid4())
    try:
        with db() as connection:
            connection.execute(
                "INSERT INTO banks(id,name,description,created_at,owner_user_id) VALUES(?,?,?,?,?)",
                (bank_id, payload.name.strip(), payload.description, int(time.time()), user_id),
            )
    except sqlite3.IntegrityError:
        raise HTTPException(409, "同名题库已存在")
    return {"id": bank_id, "message": "题库已创建"}


@app.put("/api/banks/{bank_id}")
def update_bank(bank_id: str, payload: BankIn, user_id: str = Depends(require_auth)):
    with db() as connection:
        updated = connection.execute(
            "UPDATE banks SET name=?, description=? WHERE id=? AND owner_user_id=?",
            (payload.name.strip(), payload.description, bank_id, user_id),
        ).rowcount
    if not updated:
        raise HTTPException(404, "题库不存在")
    return {"message": "题库已更新"}


@app.delete("/api/banks/{bank_id}")
def delete_bank(bank_id: str, user_id: str = Depends(require_auth)):
    with db() as connection:
        if not owned_bank(connection, bank_id, user_id):
            raise HTTPException(404, "题库不存在")
        count = connection.execute("SELECT COUNT(*) FROM questions WHERE bank_id=? AND owner_user_id=?", (bank_id, user_id)).fetchone()[0]
        if count:
            raise HTTPException(400, "题库下仍有题目，先清空后再删除")
        removed = connection.execute("DELETE FROM banks WHERE id=? AND owner_user_id=?", (bank_id, user_id)).rowcount
    if not removed:
        raise HTTPException(404, "题库不存在")
    return {"message": "题库已删除"}


@app.get("/api/stats")
def stats(user_id: str = Depends(require_auth)):
    with db() as connection:
        total = connection.execute("SELECT COUNT(*) FROM questions WHERE owner_user_id=? AND parent_id IS NULL", (user_id,)).fetchone()[0]
        chunks = connection.execute("SELECT COUNT(*) FROM questions WHERE owner_user_id=? AND parent_id IS NOT NULL", (user_id,)).fetchone()[0]
        banks_total = connection.execute("SELECT COUNT(*) FROM banks WHERE owner_user_id=?", (user_id,)).fetchone()[0]
        categories = connection.execute("SELECT category, COUNT(*) count FROM questions WHERE owner_user_id=? AND parent_id IS NULL GROUP BY category ORDER BY count DESC", (user_id,)).fetchall()
    return {"total": total, "chunks": chunks, "banks": banks_total, "categories": [dict(item) for item in categories]}


@app.get("/api/options")
def options(user_id: str = Depends(require_auth)):
    result = {}
    with db() as connection:
        for field in ("category", "difficulty", "position", "source"):
            result[field] = [item[0] for item in connection.execute(f"SELECT DISTINCT {field} FROM questions WHERE owner_user_id=? ORDER BY {field}", (user_id,)).fetchall()]
    return result


@app.get("/api/questions")
def questions(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    keyword: str = "",
    bank_id: str = "",
    user_id: str = Depends(require_auth),
):
    where, args = ["parent_id IS NULL", "owner_user_id=?"], [user_id]
    if keyword:
        where.append("(question LIKE ? OR answer LIKE ? OR tags LIKE ? OR keywords LIKE ?)")
        args += [f"%{keyword}%", f"%{keyword}%", f"%{keyword}%", f"%{keyword}%"]
    if bank_id:
        require_owned_bank(bank_id, user_id)
        where.append("bank_id=?")
        args.append(bank_id)
    clause = "WHERE " + " AND ".join(where) if where else ""
    with db() as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM questions {clause}", args).fetchone()[0]
        rows = connection.execute(
            f"SELECT * FROM questions {clause} ORDER BY created_at DESC LIMIT ? OFFSET ?",
            [*args, page_size, (page - 1) * page_size],
        ).fetchall()
    items = []
    for row in rows:
        item = serialize(row)
        item["preview"] = block_preview(item.get("answer", ""))
        items.append(item)
    return {"items": items, "total": total}


@app.post("/api/questions", status_code=201)
def create_question(payload: QuestionIn, user_id: str = Depends(require_auth)):
    try:
        ids = save_question(payload, user_id, chunk=False)
    except sqlite3.IntegrityError:
        raise HTTPException(409, "相同题目和答案已存在")
    return {"ids": ids, "message": "题目已保存"}


@app.get("/api/questions/{item_id}")
def get_question(item_id: str, user_id: str = Depends(require_auth)):
    with db() as connection:
        row = connection.execute("SELECT * FROM questions WHERE id=? AND owner_user_id=?", (item_id, user_id)).fetchone()
        if not row:
            raise HTTPException(404, "题目不存在")
        children = connection.execute("SELECT * FROM questions WHERE parent_id=? AND owner_user_id=? ORDER BY chunk_index", (item_id, user_id)).fetchall()
    item = serialize(row)
    item["children"] = [serialize(child) for child in children]
    return item


@app.put("/api/questions/{item_id}")
def update_question(item_id: str, payload: QuestionPatchIn, user_id: str = Depends(require_auth)):
    fields = []
    changed_fields = set()
    args = []
    for field in ("question", "answer", "category", "difficulty", "position", "source", "bank_id"):
        value = getattr(payload, field)
        if value is not None:
            fields.append(f"{field}=?")
            changed_fields.add(field)
            args.append(value)
    if payload.keywords is not None:
        fields.append("keywords=?")
        changed_fields.add("keywords")
        args.append(json.dumps(normalize_list(payload.keywords), ensure_ascii=False))
    if payload.tags is not None:
        fields.append("tags=?")
        changed_fields.add("tags")
        args.append(json.dumps(normalize_list(payload.tags), ensure_ascii=False))
    if not fields:
        return {"message": "没有变更"}
    if payload.bank_id is not None:
        require_owned_bank(payload.bank_id, user_id)
    args.extend([int(time.time()), item_id])
    with db() as connection:
        updated = connection.execute(
            f"UPDATE questions SET {', '.join(fields)}, updated_at=? WHERE id=? AND owner_user_id=?",
            [*args, user_id],
        ).rowcount
    if not updated:
        raise HTTPException(404, "题目不存在")
    if changed_fields & {"question", "answer", "category", "difficulty", "position", "keywords", "tags", "bank_id"}:
        refresh_question_embedding(item_id)
    return {"message": "题目已更新"}


@app.post("/api/questions/{item_id}/auto-tags")
async def auto_tags(item_id: str, user_id: str = Depends(require_auth)):
    with db() as connection:
        row = connection.execute("SELECT * FROM questions WHERE id=? AND owner_user_id=?", (item_id, user_id)).fetchone()
    if not row:
        raise HTTPException(404, "题目不存在")
    item = serialize(row)
    result = await openai_compatible_json(
        "你是题库标注助手，只输出 JSON。",
        f"请为题目生成结构化标注。\n问题：{item['question']}\n答案：{item['answer']}\n只输出 JSON，字段为 category, difficulty, position, keywords 数组, tags 数组。",
        resolve_ai_config(user_id),
    )
    patch = QuestionPatchIn(
        category=result.get("category"),
        difficulty=result.get("difficulty"),
        position=result.get("position"),
        keywords=result.get("keywords"),
        tags=result.get("tags"),
    )
    update_question(item_id, patch, user_id)
    return result


@app.post("/api/import/preview")
async def preview_import_file(
    file: UploadFile = File(...),
    chunk_mode: str = Form("smart"),
    chunk_size: int = Form(900),
    chunk_overlap: int = Form(120),
    user_id: str = Depends(require_auth),
):
    if not file.filename or not file.filename.lower().endswith((".json", ".csv", ".md", ".markdown", ".txt")):
        raise HTTPException(400, "仅支持 JSON、CSV、Markdown 和 TXT 文件")
    if chunk_mode not in {"smart", "fixed", "none"}:
        raise HTTPException(400, "分块策略无效")
    raw = await file.read()
    if len(raw) > 20 * 1024 * 1024:
        raise HTTPException(400, "题库文件不能超过 20 MB")
    try:
        parsed = parse_import(file.filename, raw)
    except Exception as error:
        raise HTTPException(400, f"文件解析失败：{error}") from error
    if not parsed:
        raise HTTPException(400, "未解析到可预览的题目，请检查文件结构")
    if len(parsed) > 1000:
        raise HTTPException(400, "单次最多预览 1000 道题目")
    items = [import_item_preview(item, index, chunk_mode, chunk_size, chunk_overlap) for index, item in enumerate(parsed, 1)]
    return {
        "filename": file.filename,
        "items": items,
        "total": len(items),
        "valid_total": sum(1 for item in items if item["valid"]),
        "estimated_chunks": sum(item.get("estimated_chunks", 0) for item in items if item["valid"]),
        "splitter": "LangChain RecursiveCharacterTextSplitter",
    }


@app.post("/api/import")
async def import_file(
    file: UploadFile = File(...),
    bank_id: str = Form(""),
    selected_indices: str = Form(""),
    chunk_mode: str = Form("smart"),
    chunk_size: int = Form(900),
    chunk_overlap: int = Form(120),
    user_id: str = Depends(require_auth),
):
    if not file.filename or not file.filename.lower().endswith((".json", ".csv", ".md", ".markdown", ".txt")):
        raise HTTPException(400, "仅支持 JSON、CSV、Markdown 和 TXT 文件")
    if chunk_mode not in {"smart", "fixed", "none"}:
        raise HTTPException(400, "分块策略无效")
    chunk_size = max(300, min(chunk_size, 4000))
    chunk_overlap = max(0, min(chunk_overlap, min(chunk_size // 2, 600)))
    raw = await file.read()
    if len(raw) > 20 * 1024 * 1024:
        raise HTTPException(400, "题库文件不能超过 20 MB")
    items = parse_import(file.filename, raw)
    if not items:
        raise HTTPException(400, "未解析到可导入题目，请检查文件结构")
    try:
        selected = {int(item) for item in json.loads(selected_indices)} if selected_indices else set(range(1, len(items) + 1))
    except (TypeError, ValueError, json.JSONDecodeError) as error:
        raise HTTPException(400, "所选题目参数无效") from error
    if not selected:
        raise HTTPException(400, "请至少选择一道题目")
    created_rows, created_questions, created_chunks, skipped, errors = 0, 0, 0, 0, []
    imported_banks = {}
    for index, item in enumerate(items, 1):
        if index not in selected:
            continue
        try:
            if not isinstance(item, dict):
                raise ValueError("该记录不是对象")
            item = dict(item)
            target_bank_id = bank_id or item.get("bank_id") or ""
            bank_name = item.pop("bank_name", "") or item.pop("bank", "") or item.pop("\u9898\u5e93", "")
            if not target_bank_id:
                target_bank_id = ensure_bank_by_name(bank_name, user_id) if bank_name else default_bank_id(user_id)
            require_owned_bank(target_bank_id, user_id)
            item["bank_id"] = target_bank_id
            ids = save_question(
                item,
                user_id,
                chunk=chunk_mode != "none",
                chunk_mode=chunk_mode,
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
            )
            created_rows += len(ids)
            created_questions += 1
            created_chunks += max(0, len(ids) - 1)
            imported_banks[target_bank_id] = imported_banks.get(target_bank_id, 0) + 1
        except sqlite3.IntegrityError:
            skipped += 1
        except Exception as exc:
            errors.append({"row": index, "message": type(exc).__name__, "detail": str(exc)})
    return {
        "filename": file.filename,
        "created": created_rows,
        "created_questions": created_questions,
        "created_chunks": created_chunks,
        "selected": len(selected),
        "skipped": skipped,
        "errors": errors,
        "banks": imported_banks,
        "splitter": "LangChain RecursiveCharacterTextSplitter",
    }


@app.delete("/api/questions/{item_id}")
def delete_question(item_id: str, user_id: str = Depends(require_auth)):
    with db() as connection:
        ids = [row["id"] for row in connection.execute("SELECT id FROM questions WHERE owner_user_id=? AND (id=? OR parent_id=?)", (user_id, item_id, item_id)).fetchall()]
        removed = connection.execute("DELETE FROM questions WHERE owner_user_id=? AND (id=? OR parent_id=?)", (user_id, item_id, item_id)).rowcount
    if not removed:
        raise HTTPException(404, "题目不存在")
    app.state.milvus.delete(COLLECTION, ids=ids)
    return {"message": "题目已删除"}


@app.post("/api/search")
def search(payload: SearchIn, user_id: str = Depends(require_auth)):
    if payload.bank_id:
        require_owned_bank(payload.bank_id, user_id)
    payload.limit = min(payload.limit, SEARCH_TOP_K)
    started = time.perf_counter()
    result = ranked_search(payload, user_id)
    latency_ms = round((time.perf_counter() - started) * 1000, 2)
    with db() as connection:
        connection.execute(
            "INSERT INTO search_logs(query,mode,result_count,created_at,user_id,latency_ms,top1_score) VALUES(?,?,?,?,?,?,?)",
            (payload.query, payload.mode, len(result), int(time.time()), user_id, latency_ms, float(result[0].get("score", 0)) if result else 0),
        )
    return {"items": result, "total": len(result), "latency_ms": latency_ms}


@app.post("/api/search/compare")
def search_compare(payload: SearchIn, user_id: str = Depends(require_auth)):
    if payload.bank_id:
        require_owned_bank(payload.bank_id, user_id)
    limit = min(payload.limit, SEARCH_TOP_K)
    return {
        "semantic": ranked_search(payload.model_copy(update={"mode": "semantic", "limit": limit}), user_id),
        "keyword": ranked_search(payload.model_copy(update={"mode": "keyword", "limit": limit}), user_id),
        "hybrid": ranked_search(payload.model_copy(update={"mode": "hybrid", "limit": limit}), user_id),
    }


@app.post("/api/ask")
async def ask(payload: AskIn, user_id: str = Depends(require_auth)):
    ai_config = resolve_ai_config(user_id)
    if not ai_config:
        raise HTTPException(400, "当前账号尚未配置并启用 DeepSeek/API，请先进入 AI 配置页面")
    if payload.bank_id:
        require_owned_bank(payload.bank_id, user_id)
    retrieval_started = time.perf_counter()
    sources = ranked_search(SearchIn(query=payload.query, bank_id=payload.bank_id, mode="hybrid", limit=SEARCH_TOP_K), user_id)
    retrieval_latency_ms = round((time.perf_counter() - retrieval_started) * 1000, 2)
    context = "\n\n".join([f"[{i + 1}] 问题：{item['question']}\n答案：{item['answer']}" for i, item in enumerate(sources)])
    generation_started = time.perf_counter()
    answer = await openai_compatible_json(
        "你是鉴微系统里的资深技术面试官。只输出 JSON。",
        f"用户问题：{payload.query}\n\nTop 3 检索资料：\n{context}\n\n请输出 JSON，字段 answer, sources_summary, confidence。",
        ai_config,
    )
    generation_latency_ms = round((time.perf_counter() - generation_started) * 1000, 2)
    with db() as connection:
        connection.execute(
            "INSERT INTO search_logs(query,mode,result_count,created_at,user_id,latency_ms,top1_score) VALUES(?,?,?,?,?,?,?)",
            (payload.query, "hybrid_ask", len(sources), int(time.time()), user_id, retrieval_latency_ms, float(sources[0].get("score", 0)) if sources else 0),
        )
    answer["metrics"] = {
        "retrieval_latency_ms": retrieval_latency_ms,
        "generation_latency_ms": generation_latency_ms,
        "total_latency_ms": round(retrieval_latency_ms + generation_latency_ms, 2),
        "source_count": len(sources),
    }
    return {"answer": answer, "sources": sources}


@app.post("/api/interviews/start")
def interview_start(payload: InterviewStartIn, user_id: str = Depends(require_auth)):
    require_owned_bank(payload.bank_id, user_id)
    with db() as connection:
        rows = connection.execute("SELECT id FROM questions WHERE bank_id=? AND owner_user_id=? AND parent_id IS NULL ORDER BY RANDOM() LIMIT 6", (payload.bank_id, user_id)).fetchall()
    if not rows:
        raise HTTPException(400, "该题库暂无题目")
    interview_id = str(uuid.uuid4())
    question_ids = [row["id"] for row in rows]
    with db() as connection:
        connection.execute(
            "INSERT INTO interviews(id,bank_id,user_id,question_ids,created_at) VALUES(?,?,?,?,?)",
            (interview_id, payload.bank_id, user_id, json.dumps(question_ids), int(time.time())),
        )
    return interview_get(interview_id, user_id)


@app.post("/api/project-interviews/start")
async def project_interview_start(
    project_title: str = Form(""),
    project_info: str = Form(""),
    bank_id: str = Form(""),
    user_id: str = Form("default"),
    file: UploadFile | None = File(default=None),
    auth_user_id: str = Depends(require_auth),
):
    document_text = ""
    if file and file.filename:
        raw = await file.read()
        if len(raw) > 10 * 1024 * 1024:
            raise HTTPException(400, "Project document must be smaller than 10 MB.")
        document_text = extract_project_document(file.filename, raw)
    context = "\n\n".join(part.strip() for part in [project_info, document_text] if part and part.strip())
    context = context[:24000]
    if len(context) < 30:
        raise HTTPException(400, "Please upload a project document or enter detailed project information.")
    title = project_title.strip() or (file.filename.rsplit(".", 1)[0] if file and file.filename else "Project")
    ai_config = resolve_ai_config(auth_user_id)
    if ai_config:
        try:
            generated, analysis = await ProjectInterviewAgent(ai_config).generate_questions(title, context)
            questions = normalize_project_questions(generated.get("questions", []), title)
            logger.info("LangChain project interview agent generated questions with %s analysis fields", len(analysis))
        except Exception as error:
            logger.exception("LangChain project interview agent failed")
            raise HTTPException(502, "项目深挖 Agent 生成失败，请稍后重试") from error
    else:
        questions = project_question_fallback(title)
    interview_id = str(uuid.uuid4())
    target_bank_id = bank_id or default_bank_id(auth_user_id)
    require_owned_bank(target_bank_id, auth_user_id)
    with db() as connection:
        connection.execute(
            """
            INSERT INTO interviews(
                id,bank_id,user_id,question_ids,created_at,mode,project_title,project_context,generated_questions
            ) VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (
                interview_id,
                target_bank_id,
                auth_user_id,
                "[]",
                int(time.time()),
                "project",
                title,
                context,
                json.dumps(questions, ensure_ascii=False),
            ),
        )
    return interview_get(interview_id, auth_user_id)


@app.get("/api/interviews/{interview_id}")
def interview_get(interview_id: str, user_id: str = Depends(require_auth)):
    with db() as connection:
        interview = connection.execute("SELECT * FROM interviews WHERE id=? AND user_id=?", (interview_id, user_id)).fetchone()
        if not interview:
            raise HTTPException(404, "面试不存在")
        ids = json.loads(interview["question_ids"] or "[]")
        questions_rows = []
        if ids:
            placeholders = ",".join(["?"] * len(ids))
            questions_rows = connection.execute(
                f"SELECT * FROM questions WHERE owner_user_id=? AND id IN ({placeholders})",
                [user_id, *ids],
            ).fetchall()
        turns = connection.execute("SELECT * FROM interview_turns WHERE interview_id=? ORDER BY created_at", (interview_id,)).fetchall()
    questions_by_id = {row["id"]: serialize(row) for row in questions_rows}
    generated_questions = safe_json_list(interview["generated_questions"] if "generated_questions" in interview.keys() else "[]")
    questions = generated_questions if interview["mode"] == "project" and generated_questions else [questions_by_id[item] for item in ids if item in questions_by_id]
    return {
        "id": interview_id,
        "bank_id": interview["bank_id"],
        "mode": interview["mode"] or "general",
        "project_title": interview["project_title"] or "",
        "questions": questions,
        "turns": [dict(row) for row in turns],
        "agent": {"framework": "LangChain", "name": "project-interview-agent"} if interview["mode"] == "project" else None,
    }


@app.post("/api/interviews/{interview_id}/answer")
async def interview_answer(interview_id: str, payload: InterviewAnswerIn, user_id: str = Depends(require_auth)):
    with db() as connection:
        interview = connection.execute("SELECT * FROM interviews WHERE id=? AND user_id=?", (interview_id, user_id)).fetchone()
        question = connection.execute("SELECT * FROM questions WHERE id=? AND owner_user_id=?", (payload.question_id, user_id)).fetchone()
    if not interview:
        raise HTTPException(404, "Interview not found.")
    item = serialize(question) if question else None
    if not item and interview["mode"] == "project":
        item = next(
            (candidate for candidate in safe_json_list(interview["generated_questions"]) if candidate.get("id") == payload.question_id),
            None,
        )
    if not item:
        raise HTTPException(404, "题目不存在")
    next_depth = payload.depth + 1
    ai_config = resolve_ai_config(user_id)
    if ai_config and interview["mode"] == "project":
        try:
            result = await ProjectInterviewAgent(ai_config).evaluate_answer(interview, item, payload)
        except Exception as error:
            logger.exception("LangChain project interview agent evaluation failed")
            raise HTTPException(502, "项目深挖 Agent 评估失败，请稍后重试") from error
    elif ai_config:
        result = await openai_compatible_json(
            "你是一名严格但友好的技术面试官。无论回答对错，都必须给出反馈，并且返回 JSON。",
            f"原题：{item['question']}\n标准答案：{item['answer']}\n本轮问题：{payload.prompt}\n候选人回答：{payload.answer}\n\n请输出 JSON，字段 score 0-100, feedback, correct_answer, strengths 数组, weaknesses 数组, follow_up 字符串。当前追问深度：{payload.depth}，最多 2。",
            ai_config,
        )
    else:
        result = {
            "score": 60 if payload.answer.strip() else 0,
            "feedback": "已收到回答，当前未启用大模型评估。",
            "correct_answer": item["answer"],
            "strengths": [],
            "weaknesses": [],
            "follow_up": "请再补充一下这个题目的关键边界条件。",
        }
    follow_up = result.get("follow_up") if next_depth <= 2 else ""
    turn_id = str(uuid.uuid4())
    score = float(result.get("score", 0) or 0)
    now = int(time.time())
    with db() as connection:
        connection.execute(
            "INSERT INTO interview_turns VALUES(?,?,?,?,?,?,?,?,?,?)",
            (
                turn_id,
                interview_id,
                payload.question_id,
                payload.prompt,
                payload.answer,
                json.dumps(result, ensure_ascii=False),
                follow_up,
                payload.depth,
                score,
                now,
            ),
        )
        if question:
            next_review_at = now if score < 70 else now + 86400
            connection.execute(
                """
                INSERT INTO user_question_states(user_id,question_id,mastery_level,next_review_at,last_rating)
                VALUES(?,?,?,?,?)
                ON CONFLICT(user_id,question_id) DO UPDATE SET
                  mastery_level=excluded.mastery_level,
                  next_review_at=MIN(COALESCE(user_question_states.next_review_at,excluded.next_review_at),excluded.next_review_at),
                  last_rating=excluded.last_rating
                """,
                (user_id, payload.question_id, round(score / 100, 3), next_review_at, "interview"),
            )
    return {"id": turn_id, "evaluation": result, "follow_up": follow_up, "next_depth": next_depth}


@app.post("/api/interview-turns/{turn_id}/save-question")
def save_turn_question(turn_id: str, bank_id: str = "", user_id: str = Depends(require_auth)):
    with db() as connection:
        turn = connection.execute("SELECT * FROM interview_turns WHERE id=?", (turn_id,)).fetchone()
        original = connection.execute("SELECT * FROM questions WHERE id=? AND owner_user_id=?", (turn["question_id"], user_id)).fetchone() if turn else None
        interview = connection.execute("SELECT * FROM interviews WHERE id=? AND user_id=?", (turn["interview_id"], user_id)).fetchone() if turn else None
    if not turn or not turn["follow_up"]:
        raise HTTPException(404, "追问不存在")
    original_item = serialize(original) if original else None
    if not original_item and interview and interview["mode"] == "project":
        original_item = next(
            (item for item in safe_json_list(interview["generated_questions"]) if item.get("id") == turn["question_id"]),
            None,
        )
    if not original_item:
        raise HTTPException(404, "Original interview question not found.")
    target_bank_id = bank_id or original_item.get("bank_id") or (interview["bank_id"] if interview else "")
    require_owned_bank(target_bank_id, user_id)
    ids = save_question(
        {
            "question": turn["follow_up"],
            "answer": "由模拟面试追问生成，请补充标准答案。",
            "category": original_item["category"],
            "difficulty": original_item["difficulty"],
            "position": original_item["position"],
            "keywords": original_item["keywords"],
            "tags": ["AI追问"],
            "source": "模拟面试追问",
            "bank_id": target_bank_id,
        },
        user_id,
        chunk=False,
    )
    return {"ids": ids, "message": "追问已加入题库"}


@app.post("/api/interview-turns/{turn_id}/save-current-question")
def save_current_interview_question(turn_id: str, bank_id: str = "", user_id: str = Depends(require_auth)):
    with db() as connection:
        turn = connection.execute("SELECT * FROM interview_turns WHERE id=?", (turn_id,)).fetchone()
        interview = connection.execute("SELECT * FROM interviews WHERE id=? AND user_id=?", (turn["interview_id"], user_id)).fetchone() if turn else None
        original = connection.execute("SELECT * FROM questions WHERE id=? AND owner_user_id=?", (turn["question_id"], user_id)).fetchone() if turn else None
    if not turn or not interview:
        raise HTTPException(404, "Interview turn not found.")
    item = serialize(original) if original else next(
        (candidate for candidate in safe_json_list(interview["generated_questions"]) if candidate.get("id") == turn["question_id"]),
        None,
    )
    if not item:
        raise HTTPException(404, "Interview question not found.")
    try:
        evaluation = json.loads(turn["feedback"] or "{}")
    except Exception:
        evaluation = {}
    target_bank_id = bank_id or interview["bank_id"]
    require_owned_bank(target_bank_id, user_id)
    ids = save_question(
        {
            "question": turn["prompt"] or item["question"],
            "answer": evaluation.get("correct_answer") or item.get("answer") or "",
            "category": item.get("category") or "\u9879\u76ee\u6df1\u6316",
            "difficulty": item.get("difficulty") or "\u56f0\u96be",
            "position": item.get("position") or "\u901a\u7528",
            "keywords": item.get("keywords", []),
            "tags": list(dict.fromkeys(normalize_list(item.get("tags", [])) + ["\u9762\u8bd5\u6536\u85cf"])),
            "source": "\u9879\u76ee\u6df1\u6316\u9762\u8bd5" if interview["mode"] == "project" else "\u901a\u7528\u9898\u5e93\u9762\u8bd5",
            "bank_id": target_bank_id,
        },
        user_id,
        chunk=False,
    )
    return {"ids": ids, "message": "Question saved to bank."}


@app.get("/api/interviews/{interview_id}/report")
def interview_report(interview_id: str, user_id: str = Depends(require_auth)):
    with db() as connection:
        interview = connection.execute("SELECT * FROM interviews WHERE id=? AND user_id=?", (interview_id, user_id)).fetchone()
        turns = connection.execute("SELECT * FROM interview_turns WHERE interview_id=? ORDER BY created_at", (interview_id,)).fetchall()
    if not interview:
        raise HTTPException(404, "Interview not found.")
    evaluations = []
    for turn in turns:
        try:
            evaluation = json.loads(turn["feedback"] or "{}")
        except Exception:
            evaluation = {}
        if evaluation:
            evaluations.append(evaluation)
    scores = [float(item.get("score", 0) or 0) for item in evaluations]
    dimension_values = {}
    strengths = []
    weaknesses = []
    for evaluation in evaluations:
        strengths.extend(str(item) for item in evaluation.get("strengths", []) if str(item).strip())
        weaknesses.extend(str(item) for item in evaluation.get("weaknesses", []) if str(item).strip())
        for key, value in (evaluation.get("dimensions") or {}).items():
            try:
                dimension_values.setdefault(key, []).append(float(value))
            except (TypeError, ValueError):
                continue
    dimensions = {key: round(sum(values) / len(values), 1) for key, values in dimension_values.items() if values}
    report_score = round(sum(scores) / len(scores), 1) if scores else 0
    completed_at = int(time.time())
    with db() as connection:
        connection.execute(
            "UPDATE interviews SET completed_at=?,report_score=? WHERE id=? AND user_id=?",
            (completed_at, report_score, interview_id, user_id),
        )
        complete_next_interview_task(connection, user_id, interview["mode"] or "general", interview_id)
    return {
        "interview_id": interview_id,
        "mode": interview["mode"] or "general",
        "project_title": interview["project_title"] or "",
        "answered_turns": len(evaluations),
        "score": report_score,
        "dimensions": dimensions,
        "strengths": list(dict.fromkeys(strengths))[:6],
        "weaknesses": list(dict.fromkeys(weaknesses))[:6],
    }


@app.get("/api/reviews")
def reviews(user_id: str = "default", scope: Literal["today", "yesterday"] = "today", limit: int = 30, auth_user_id: str = Depends(require_auth)):
    user_id = auth_user_id
    now = int(time.time())
    day = 86400
    with db() as connection:
        if scope == "yesterday":
            start = now - day * 2
            end = now - day
            rows = connection.execute(
                """
                SELECT DISTINCT q.*,s.mastery_level,s.review_count,s.last_reviewed_at,
                s.next_review_at,s.interval_days,s.ease_factor,s.stability,s.difficulty_factor,s.lapse_count,s.last_rating
                FROM questions q
                JOIN user_question_states s ON s.question_id=q.id AND s.user_id=?
                JOIN interview_turns t ON t.question_id=q.id
                WHERE q.owner_user_id=? AND s.last_reviewed_at BETWEEN ? AND ?
                ORDER BY s.last_reviewed_at DESC LIMIT ?
                """,
                (user_id, user_id, start, end, limit),
            ).fetchall()
        else:
            rows = connection.execute(
                """
                SELECT DISTINCT q.*,s.mastery_level,s.review_count,s.last_reviewed_at,
                s.next_review_at,s.interval_days,s.ease_factor,s.stability,s.difficulty_factor,s.lapse_count,s.last_rating
                FROM questions q
                JOIN interview_turns t ON t.question_id=q.id
                LEFT JOIN user_question_states s ON s.question_id=q.id AND s.user_id=?
                WHERE q.owner_user_id=? AND (s.next_review_at IS NULL OR s.next_review_at<=?)
                ORDER BY COALESCE(s.next_review_at,0),q.created_at LIMIT ?
                """,
                (user_id, user_id, now, limit),
            ).fetchall()
    items = []
    for row in rows:
        item = serialize(row)
        item["recall_probability"] = recall_probability(item.get("stability"), item.get("last_reviewed_at"))
        items.append(item)
    return {"items": items, "total": len(items)}


@app.post("/api/reviews/{item_id}")
def review(item_id: str, payload: ReviewIn, user_id: str = Depends(require_auth)):
    now = int(time.time())
    weights = {"again": 0, "hard": 1, "good": 2, "easy": 3}
    with db() as connection:
        if not connection.execute("SELECT 1 FROM questions WHERE id=? AND owner_user_id=?", (item_id, user_id)).fetchone():
            raise HTTPException(404, "题目不存在")
        row = connection.execute("SELECT * FROM user_question_states WHERE user_id=? AND question_id=?", (user_id, item_id)).fetchone()
        count = (row["review_count"] if row else 0) + 1
        stability = float(row["stability"] if row else 1.0)
        difficulty = float(row["difficulty_factor"] if row else 5.0)
        lapse_count = int(row["lapse_count"] if row else 0)
        if payload.rating == "again":
            stability = max(0.5, stability * 0.45)
            difficulty = min(10, difficulty + 1.2)
            lapse_count += 1
            interval = 1
        elif payload.rating == "hard":
            stability = max(1.0, stability * 1.35)
            difficulty = min(10, difficulty + 0.35)
            interval = max(1, round(stability))
        elif payload.rating == "good":
            stability = stability * (2.2 + (10 - difficulty) / 10)
            difficulty = max(1, difficulty - 0.15)
            interval = max(3, round(stability))
        else:
            stability = stability * (3.0 + (10 - difficulty) / 8)
            difficulty = max(1, difficulty - 0.45)
            interval = max(5, round(stability))
        next_at = now + interval * 86400
        connection.execute(
            """
            INSERT INTO user_question_states(user_id,question_id,mastery_level,review_count,
            last_reviewed_at,next_review_at,interval_days,ease_factor,stability,difficulty_factor,lapse_count,last_rating)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(user_id,question_id) DO UPDATE SET mastery_level=excluded.mastery_level,
            review_count=excluded.review_count,last_reviewed_at=excluded.last_reviewed_at,
            next_review_at=excluded.next_review_at,interval_days=excluded.interval_days,ease_factor=excluded.ease_factor,
            stability=excluded.stability,difficulty_factor=excluded.difficulty_factor,lapse_count=excluded.lapse_count,
            last_rating=excluded.last_rating
            """,
            (user_id, item_id, weights[payload.rating] / 3, count, now, next_at, interval, 2.5, stability, difficulty, lapse_count, payload.rating),
        )
        remaining_due = connection.execute(
            "SELECT COUNT(*) FROM user_question_states WHERE user_id=? AND (next_review_at IS NULL OR next_review_at<=?)",
            (user_id, now),
        ).fetchone()[0]
        if remaining_due == 0:
            task = connection.execute(
                """
                SELECT id FROM training_tasks
                WHERE user_id=? AND status='pending' AND task_type='review' AND due_date<=?
                ORDER BY due_date,created_at LIMIT 1
                """,
                (user_id, coach_today().isoformat()),
            ).fetchone()
            if task:
                connection.execute(
                    "UPDATE training_tasks SET status='completed',completed_at=? WHERE id=?",
                    (now, task["id"]),
                )
    return {"message": "复习记录已保存", "interval_days": interval, "next_review_at": next_at, "stability": round(stability, 2), "difficulty_factor": round(difficulty, 2), "recall_probability": 1.0}
